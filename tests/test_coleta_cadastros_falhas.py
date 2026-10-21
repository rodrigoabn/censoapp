"""Falhas de coleta de cadastros (login, escola não encontrada, erro por escola)."""

from __future__ import annotations

import io

import openpyxl
import pytest

from apps.coleta_cadastros_escolas.src import exporter, scraper

ESCOLAS = [
    {"codigo_inep": "111", "nome_unidade": "Escola Um"},
    {"codigo_inep": "222", "nome_unidade": "Escola Dois"},
]


class _LocatorFalso:
    def is_visible(self):
        return False


class _PaginaFalsa:
    def locator(self, *a, **k):
        return _LocatorFalso()

    def close(self):
        pass


class _ContextoFalso:
    def new_page(self):
        return _PaginaFalsa()

    def clear_cookies(self):
        pass

    def close(self):
        pass


class _BrowserFalso:
    def new_context(self, **kwargs):
        return _ContextoFalso()


@pytest.fixture(autouse=True)
def _sem_espera(monkeypatch):
    monkeypatch.setattr("time.sleep", lambda *_: None)


def test_login_falho_aborta_a_coleta(monkeypatch):
    def _falha(*a, **k):
        raise scraper.LoginFalhou("Nome de usuário ou senha inválida.")

    monkeypatch.setattr(scraper, "_login_and_read", _falha)
    with pytest.raises(scraper.LoginFalhou) as excinfo:
        scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS)
    assert excinfo.value.mensagem == "Nome de usuário ou senha inválida."


def test_escola_nao_encontrada_vira_falha(monkeypatch):
    monkeypatch.setattr(scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(scraper, "_session_expirada", lambda page: False)
    monkeypatch.setattr(scraper, "scrape_school", lambda *a, **k: None)
    outputs, failures = scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS[:1])
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Escola não encontrada"}
    ]


def test_escola_ok_e_erro_inesperado_convivem(monkeypatch):
    monkeypatch.setattr(scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(scraper, "_session_expirada", lambda page: False)

    def _scrape(page, codigo, nome):
        if codigo == "222":
            raise RuntimeError("boom")
        return {"Vinculação institucional e Convênio": {"Código do Inep": codigo}}

    monkeypatch.setattr(scraper, "scrape_school", _scrape)
    outputs, failures = scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS)
    assert [e["codigo_inep"] for e, _ in outputs] == ["111"]
    assert failures == [
        {"codigo_inep": "222", "nome_unidade": "Escola Dois", "motivo": "Erro: boom", "timeout": False}
    ]


def test_planilha_de_nao_coletadas_tem_cabecalho_e_linhas():
    data = exporter.build_pending_xlsx([
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Falha no login", "timeout": False},
    ])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["Código do Inep", "Nome da Unidade", "Motivo"]
    assert [c.value for c in ws[2]] == ["111", "Escola Um", "Falha no login"]
    assert exporter.pending_filename().startswith("cadastro_escola_escolas_nao_coletadas_")


def test_re_login_falho_apos_sessao_expirada(monkeypatch):
    chamadas = {"login": 0}

    def _login(page, login, senha):
        chamadas["login"] += 1
        if chamadas["login"] > 1:
            raise scraper.LoginFalhou()

    monkeypatch.setattr(scraper, "_login_and_read", _login)
    monkeypatch.setattr(scraper, "_session_expirada", lambda page: False)

    def _scrape(page, codigo, nome):
        raise scraper.SessaoExpirada()

    monkeypatch.setattr(scraper, "scrape_school", _scrape)
    outputs, failures = scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS[:1])
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Erro (re-login): LoginFalhou", "timeout": False}
    ]


def test_re_login_falho_na_checagem_previa(monkeypatch):
    chamadas = {"login": 0}

    def _login(page, login, senha):
        chamadas["login"] += 1
        if chamadas["login"] > 1:
            raise scraper.LoginFalhou()

    monkeypatch.setattr(scraper, "_login_and_read", _login)
    monkeypatch.setattr(scraper, "_session_expirada", lambda page: True)
    outputs, failures = scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS[:1])
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Erro (re-login): LoginFalhou", "timeout": False}
    ]


def test_planilha_de_nao_coletadas_neutraliza_formulas():
    data = exporter.build_pending_xlsx([
        {"codigo_inep": "111", "nome_unidade": "=1+1", "motivo": "Falha no login", "timeout": False},
    ])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["B2"].value == "'=1+1"


def test_interrupcao_antes_da_primeira_escola_marca_todas_como_nao_processadas(monkeypatch):
    monkeypatch.setattr(scraper, "_login_and_read", lambda *a, **k: None)
    outputs, failures = scraper.run_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, stop_check=lambda: True
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Coleta interrompida (não processada)",
        },
        {
            "codigo_inep": "222",
            "nome_unidade": "Escola Dois",
            "motivo": "Coleta interrompida (não processada)",
        },
    ]


def test_interrupcao_apos_coletar_descarta_a_atual_e_marca_a_restante(monkeypatch):
    monkeypatch.setattr(scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(scraper, "_session_expirada", lambda page: False)
    monkeypatch.setattr(
        scraper, "scrape_school", lambda page, codigo, nome: {"Ficha": {"Código do Inep": codigo}}
    )

    chamadas = {"n": 0}

    def _stop():
        chamadas["n"] += 1
        return chamadas["n"] > 3  # True só depois de coletar a 2ª escola

    outputs, failures = scraper.run_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, stop_check=_stop
    )
    assert [e["codigo_inep"] for e, _ in outputs] == ["111"]
    assert failures == [
        {
            "codigo_inep": "222",
            "nome_unidade": "Escola Dois",
            "motivo": "Coleta interrompida (não processada)",
        }
    ]
