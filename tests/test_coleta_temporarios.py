"""Temporários exclusivos por execução — regressão do vazamento entre execuções
concorrentes (antes, os 4 coletores usavam uma pasta fixa em /tmp)."""

from __future__ import annotations

import io
import zipfile

import openpyxl

from apps.coleta_cadastros_escolas.src import exporter as cad_exporter
from apps.coleta_cadastros_escolas.src import modelo as cad_modelo
from apps.coleta_gestor.src import exporter as ges_exporter
from apps.coleta_gestor.src import modelo as ges_modelo
from apps.coleta_recibos_fechamento.src import zipper as rec_zipper
from apps.coleta_relatorios.src import zipper as rel_zipper


def test_pastas_dos_quatro_coletores_sao_unicas_por_execucao():
    pastas = [
        cad_exporter.nova_pasta_temporaria(),
        ges_exporter.nova_pasta_temporaria(),
        rel_zipper.nova_pasta_temporaria(),
        rec_zipper.nova_pasta_temporaria(),
    ]
    try:
        assert len({str(p) for p in pastas}) == 4
        assert all(p.is_dir() for p in pastas)
        prefixos = [
            "educacenso_tmp_",
            "educacenso_gestor_tmp_",
            "educacenso_relatorios_tmp_",
            "educacenso_recibos_fechamento_tmp_",
        ]
        for p, pref in zip(pastas, prefixos):
            assert p.name.startswith(pref)
    finally:
        for mod, p in zip([cad_exporter, ges_exporter, rel_zipper, rec_zipper], pastas):
            mod.remover_temp(p)


def test_cadastros_grava_e_consolida_na_pasta_da_execucao():
    pasta = cad_exporter.nova_pasta_temporaria()
    try:
        model = cad_modelo.load_model()
        sheet = next(iter(model))
        dados = {sheet: {c: "" for c in model[sheet]}}
        dados[sheet]["Código do Inep"] = "11111111"
        dados[sheet]["Unidade Escolar"] = "Escola Um"
        cad_exporter.append_school_temp(dados, pasta)
        data = cad_exporter.build_xlsx(tmp_dir=pasta)
    finally:
        cad_exporter.remover_temp(pasta)

    ws = openpyxl.load_workbook(io.BytesIO(data))[cad_exporter._safe_sheet_title(sheet)]
    valores = [c.value for row in ws.iter_rows() for c in row]
    assert "11111111" in valores and "Escola Um" in valores
    assert not pasta.exists()  # build_xlsx remove a pasta da execução


def test_gestor_grava_e_consolida_na_pasta_da_execucao():
    pasta = ges_exporter.nova_pasta_temporaria()
    try:
        dados = {
            ges_modelo.GESTOR_SHEET: {
                c: "" for c in ges_modelo.load_model()[ges_modelo.GESTOR_SHEET]
            }
        }
        dados[ges_modelo.GESTOR_SHEET]["Código do Inep"] = "22222222"
        dados[ges_modelo.GESTOR_SHEET]["Unidade Escolar"] = "Escola Dois"
        ges_exporter.append_school_temp(dados, pasta)
        data = ges_exporter.build_xlsx(tmp_dir=pasta)
    finally:
        ges_exporter.remover_temp(pasta)

    ws = openpyxl.load_workbook(io.BytesIO(data))[ges_modelo.GESTOR_SHEET]
    valores = [c.value for row in ws.iter_rows() for c in row]
    assert "22222222" in valores and "Escola Dois" in valores
    assert not pasta.exists()


def test_zip_de_relatorios_isola_e_limpa_pasta():
    pasta = rel_zipper.nova_pasta_temporaria()
    rel_zipper.append_csv("1", "Escola Um", b"a;b\n1;2", pasta)
    assert rel_zipper.count_temp(pasta) == 1
    data = rel_zipper.build_zip(tmp_dir=pasta)
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["1_Escola Um.csv"]
    assert not pasta.exists()


def test_zip_de_recibos_grava_na_pasta_da_execucao():
    pasta = rec_zipper.nova_pasta_temporaria()
    rec_zipper.append_pdf("1", "Escola Um", b"%PDF-1.4\n%%EOF", pasta)
    data = rec_zipper.build_zip(tmp_dir=pasta)
    assert zipfile.ZipFile(io.BytesIO(data)).namelist() == ["1_Escola Um.pdf"]
    assert not pasta.exists()


def test_collectors_de_coleta_importam_sem_erro():
    import apps.coleta_gestor.src.collector  # noqa: F401
    import apps.coleta_recibos_fechamento.src.collector  # noqa: F401
    import apps.coleta_relatorios.src.collector  # noqa: F401
    from apps.coleta_cadastros_escolas.src import collector as cadastros_collector  # noqa: F401
