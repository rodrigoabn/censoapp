"""Relatórios/Recibos: falha de login total e sessão expirada durante o download."""

from __future__ import annotations

import io

import openpyxl
import pytest
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

from apps.coleta_recibos_fechamento.src import receipt_scraper
from apps.coleta_recibos_fechamento.src import zipper as rec_zipper
from apps.coleta_relatorios.src import report_scraper
from apps.coleta_relatorios.src import zipper as rel_zipper

ESCOLAS = [{"codigo_inep": "111", "nome_unidade": "Escola Um"}]


class _LocatorFalso:
    @property
    def first(self):
        return self

    def count(self):
        return 0

    def is_visible(self):
        return False


class _BotaoFalso:
    @property
    def first(self):
        return self

    def click(self, *a, **k):
        raise PlaywrightTimeoutError("timeout")


class _PaginaFalsa:
    def goto(self, *a, **k):
        return None

    def wait_for_selector(self, *a, **k):
        return None

    def wait_for_timeout(self, *a, **k):
        return None

    def wait_for_load_state(self, *a, **k):
        return None

    def locator(self, *a, **k):
        return _LocatorFalso()

    def get_by_role(self, *a, **k):
        return _BotaoFalso()

    def close(self):
        return None


class _ContextoFalso:
    def new_page(self):
        return _PaginaFalsa()

    def clear_cookies(self):
        return None

    def close(self):
        return None


class _BrowserFalso:
    def new_context(self, **kwargs):
        return _ContextoFalso()


@pytest.fixture(autouse=True)
def _sem_espera(monkeypatch):
    monkeypatch.setattr("time.sleep", lambda *_: None)
    monkeypatch.setattr(report_scraper, "_wait_loader", lambda *a, **k: None)
    monkeypatch.setattr(receipt_scraper, "_wait_loader", lambda *a, **k: None)


def test_login_falho_vira_falhas_em_relatorios(monkeypatch):
    def _falha(*a, **k):
        raise report_scraper.LoginFalhou()

    monkeypatch.setattr(report_scraper, "_login_and_read", _falha)
    outputs, failures = report_scraper.run_report_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/relatorio"
    )
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Falha no login"}
    ]


def test_login_falho_vira_falhas_em_recibos(monkeypatch):
    def _falha(*a, **k):
        raise receipt_scraper.LoginFalhou()

    monkeypatch.setattr(receipt_scraper, "_login_and_read", _falha)
    outputs, failures = receipt_scraper.run_receipt_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/fechamento"
    )
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Falha no login"}
    ]


def test_download_de_relatorio_com_sessao_expirada_levanta_sessao_expirada(monkeypatch):
    monkeypatch.setattr(report_scraper, "_session_expirada", lambda page: True)
    with pytest.raises(report_scraper.SessaoExpirada):
        report_scraper._download_report_csv(_PaginaFalsa(), "https://exemplo/relatorio")


def test_download_de_recibo_com_sessao_expirada_levanta_sessao_expirada(monkeypatch):
    monkeypatch.setattr(receipt_scraper, "_session_expirada", lambda page: True)
    with pytest.raises(receipt_scraper.SessaoExpirada):
        receipt_scraper._download_receipt_pdf(
            _ContextoFalso(), _PaginaFalsa(), "https://exemplo/fechamento"
        )


def test_relatorio_sem_preenchimento_segue_levantando_preenchimento_nao_iniciado(monkeypatch):
    monkeypatch.setattr(report_scraper, "_session_expirada", lambda page: False)
    with pytest.raises(report_scraper.PreenchimentoNaoIniciado):
        report_scraper._download_report_csv(_PaginaFalsa(), "https://exemplo/relatorio")


def test_recibo_indisponivel_segue_quando_sessao_ok(monkeypatch):
    monkeypatch.setattr(receipt_scraper, "_session_expirada", lambda page: False)
    with pytest.raises(receipt_scraper.ReciboIndisponivel):
        receipt_scraper._download_receipt_pdf(
            _ContextoFalso(), _PaginaFalsa(), "https://exemplo/fechamento"
        )


def test_planilha_de_nao_coletadas_de_relatorios_neutraliza_formulas():
    data = rel_zipper.build_pending_xlsx([
        {"codigo_inep": "111", "nome_unidade": "=1+1", "motivo": "Falha no login"},
    ])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["B2"].value == "'=1+1"


def test_planilha_de_nao_coletadas_de_recibos_neutraliza_formulas():
    data = rec_zipper.build_pending_xlsx([
        {"codigo_inep": "111", "nome_unidade": "=1+1", "motivo": "Falha no login"},
    ])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["B2"].value == "'=1+1"


DUAS_ESCOLAS = [
    {"codigo_inep": "111", "nome_unidade": "Escola Um"},
    {"codigo_inep": "222", "nome_unidade": "Escola Dois"},
]


def test_interrupcao_em_relatorios_marca_as_escolas_restantes(monkeypatch):
    monkeypatch.setattr(report_scraper, "_login_and_read", lambda *a, **k: None)
    outputs, failures = report_scraper.run_report_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/relatorio", stop_check=lambda: True
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Coleta interrompida (não processada)",
        }
    ]


def test_interrupcao_em_recibos_marca_as_escolas_restantes(monkeypatch):
    monkeypatch.setattr(receipt_scraper, "_login_and_read", lambda *a, **k: None)
    outputs, failures = receipt_scraper.run_receipt_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/fechamento", stop_check=lambda: True
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Coleta interrompida (não processada)",
        }
    ]


def test_interrupcao_apos_baixar_relatorio_descarta_a_atual_e_marca_a_restante(monkeypatch):
    monkeypatch.setattr(report_scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(report_scraper, "_session_expirada", lambda page: False)
    monkeypatch.setattr(report_scraper, "_select_school", lambda page, codigo: True)
    monkeypatch.setattr(report_scraper, "_download_report_csv", lambda page, url: b"csv")

    chamadas = {"n": 0}

    def _stop():
        chamadas["n"] += 1
        return chamadas["n"] > 3  # True só depois de baixar o relatório da 2ª escola

    outputs, failures = report_scraper.run_report_scraping(
        _BrowserFalso(), "l", "s", DUAS_ESCOLAS, "https://exemplo/relatorio", stop_check=_stop
    )
    assert [e["codigo_inep"] for e, _ in outputs] == ["111"]
    assert failures == [
        {
            "codigo_inep": "222",
            "nome_unidade": "Escola Dois",
            "motivo": "Coleta interrompida (não processada)",
        }
    ]


def test_interrupcao_apos_baixar_recibo_descarta_a_atual_e_marca_a_restante(monkeypatch):
    monkeypatch.setattr(receipt_scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(receipt_scraper, "_session_expirada", lambda page: False)
    monkeypatch.setattr(receipt_scraper, "_select_school", lambda page, codigo: True)
    monkeypatch.setattr(receipt_scraper, "_download_receipt_pdf", lambda *a, **k: b"pdf")

    chamadas = {"n": 0}

    def _stop():
        chamadas["n"] += 1
        return chamadas["n"] > 3  # True só depois de baixar o recibo da 2ª escola

    outputs, failures = receipt_scraper.run_receipt_scraping(
        _BrowserFalso(), "l", "s", DUAS_ESCOLAS, "https://exemplo/fechamento", stop_check=_stop
    )
    assert [e["codigo_inep"] for e, _ in outputs] == ["111"]
    assert failures == [
        {
            "codigo_inep": "222",
            "nome_unidade": "Escola Dois",
            "motivo": "Coleta interrompida (não processada)",
        }
    ]


# ── mapa exceção→motivo no nível de run_* (regressão Fase 4.9) ───────────────
# O loop genérico converteu exceções específicas de rel/rec em "Erro: " (e o
# log perdia o texto). Estes testes pinam os motivos/progressos do baseline.


def _progressos():
    eventos = []
    return eventos, lambda *ev: eventos.append(ev)


def _preparar_coleta(monkeypatch, modulo):
    monkeypatch.setattr(modulo, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(modulo, "_session_expirada", lambda page: False)
    monkeypatch.setattr(modulo, "_select_school", lambda page, codigo: True)


def test_preenchimento_nao_iniciado_vira_motivo_proprio_em_relatorios(monkeypatch):
    _preparar_coleta(monkeypatch, report_scraper)

    def _falha(page, url):
        raise report_scraper.PreenchimentoNaoIniciado()

    monkeypatch.setattr(report_scraper, "_download_report_csv", _falha)
    eventos, on_progress = _progressos()
    outputs, failures = report_scraper.run_report_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/relatorio", on_progress=on_progress
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Preenchimento não iniciado",
        }
    ]
    assert eventos[-1] == (1, 1, "111", "Escola Um", "Preenchimento não iniciado.")


def test_relatorio_inexistente_vira_motivo_proprio_em_relatorios(monkeypatch):
    _preparar_coleta(monkeypatch, report_scraper)

    def _falha(page, url):
        raise report_scraper.RelatorioInexistente()

    monkeypatch.setattr(report_scraper, "_download_report_csv", _falha)
    eventos, on_progress = _progressos()
    outputs, failures = report_scraper.run_report_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/relatorio", on_progress=on_progress
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Relatório inexistente",
        }
    ]
    assert eventos[-1] == (1, 1, "111", "Escola Um", "Relatório inexistente.")


def test_recibo_indisponivel_vira_motivo_proprio_em_recibos(monkeypatch):
    _preparar_coleta(monkeypatch, receipt_scraper)

    def _falha(context, page, url):
        raise receipt_scraper.ReciboIndisponivel()

    monkeypatch.setattr(receipt_scraper, "_download_receipt_pdf", _falha)
    eventos, on_progress = _progressos()
    outputs, failures = receipt_scraper.run_receipt_scraping(
        _BrowserFalso(), "l", "s", ESCOLAS, "https://exemplo/fechamento", on_progress=on_progress
    )
    assert outputs == []
    assert failures == [
        {
            "codigo_inep": "111",
            "nome_unidade": "Escola Um",
            "motivo": "Recibo de fechamento indisponível",
        }
    ]
    assert eventos[-1] == (1, 1, "111", "Escola Um", "Recibo de fechamento indisponível.")
