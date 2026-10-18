"""_finalize dos 4 frontends de coleta: arquivo principal, planilha de não
coletadas e limpeza dos temporários.

Cobre o comportamento visível ao usuário que só rodava de verdade na E2E:
quando há falhas, além do arquivo principal é oferecida a planilha "escolas
não coletadas"; quando não há coleta nenhuma, os temporários são removidos e a
pendente continua sendo oferecida (as flags `*_finalized`/`*_pending` são o que
o bloco de persistência do render usa para reexibir o botão após o rerun).
"""

from __future__ import annotations

import io
import threading
import zipfile

import openpyxl

from apps.coleta_gestor.src import exporter as ges_exporter
from apps.coleta_gestor.src import frontend as ges_frontend
from apps.coleta_gestor.src import modelo as ges_modelo
from apps.coleta_recibos_fechamento.src import frontend as rec_frontend
from apps.coleta_recibos_fechamento.src import zipper as rec_zipper
from apps.coleta_relatorios.src import frontend as rel_frontend
from apps.coleta_relatorios.src import zipper as rel_zipper
from src import exporter as cad_exporter
from src import frontend as cad_frontend
from src import modelo as cad_modelo

FALHAS = [{"codigo_inep": "222", "nome_unidade": "Escola Dois", "motivo": "Falha no login"}]


class _EstadoFalso:
    """session_state com acesso por atributo (cadastros/gestor) e por chave (relatórios/recibos)."""

    def __init__(self, **iniciais):
        object.__setattr__(self, "_dados", dict(iniciais))

    def __getattr__(self, nome):
        try:
            return self._dados[nome]
        except KeyError:
            raise AttributeError(nome) from None

    def __setattr__(self, nome, valor):
        self._dados[nome] = valor

    def __getitem__(self, chave):
        return self._dados[chave]

    def __setitem__(self, chave, valor):
        self._dados[chave] = valor


class _StFalso:
    def __init__(self, estado):
        self.session_state = estado
        self.erros, self.sucessos, self.markdowns, self.downloads = [], [], [], []

    def error(self, msg):
        self.erros.append(msg)

    def success(self, msg):
        self.sucessos.append(msg)

    def markdown(self, body, **kwargs):
        self.markdowns.append(body)

    def download_button(self, **kwargs):
        self.downloads.append(kwargs)


def test_finalize_cadastros_gera_xlsx_e_pendentes(monkeypatch):
    pasta = cad_exporter.nova_pasta_temporaria()
    modelo = cad_modelo.load_model()
    sheet = next(iter(modelo))
    dados = {sheet: {c: "" for c in modelo[sheet]}}
    dados[sheet]["Código do Inep"] = "11111111"
    dados[sheet]["Unidade Escolar"] = "Escola Um"
    cad_exporter.append_school_temp(dados, pasta)

    estado = _EstadoFalso(
        ws_stop=None,
        ws_lock=threading.Lock(),
        ws_shared={"error": None, "failures": FALHAS},
        ws_tmp_dir=pasta,
        ws_running=True,
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(cad_frontend, "st", st_falso)

    cad_frontend._finalize(1)

    assert estado.ws_running is False
    assert estado.ws_finalized is True

    xlsx, nome = estado.ws_xlsx
    assert nome.endswith(".xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))[cad_exporter._safe_sheet_title(sheet)]
    assert "11111111" in [c.value for row in ws.iter_rows() for c in row]
    assert not pasta.exists()  # build_xlsx remove a pasta da execução

    pend, pend_nome = estado.ws_pending
    assert pend_nome.startswith("cadastro_escola_escolas_nao_coletadas_")
    ws_pend = openpyxl.load_workbook(io.BytesIO(pend)).active
    assert [c.value for c in ws_pend[2]] == ["222", "Escola Dois", "Falha no login"]

    assert [d["label"] for d in st_falso.downloads] == [
        "Baixar XLSX",
        "Baixar escolas não coletadas (XLSX)",
    ]


def test_finalize_cadastros_sem_coleta_limpa_temp_e_mantem_pendentes(monkeypatch):
    pasta = cad_exporter.nova_pasta_temporaria()  # vazia: nada foi coletado

    estado = _EstadoFalso(
        ws_stop=None,
        ws_lock=threading.Lock(),
        ws_shared={"error": None, "failures": FALHAS},
        ws_tmp_dir=pasta,
        ws_running=True,
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(cad_frontend, "st", st_falso)

    cad_frontend._finalize(0)

    assert estado.ws_xlsx is None
    assert not pasta.exists()  # remover_temp limpa a pasta da execução
    assert estado.ws_finalized is True and estado.ws_pending is not None
    assert any("Arquivo XLSX não gerado" in m for m in st_falso.markdowns)
    assert [d["label"] for d in st_falso.downloads] == ["Baixar escolas não coletadas (XLSX)"]


def test_finalize_gestor_gera_xlsx_e_pendentes(monkeypatch):
    pasta = ges_exporter.nova_pasta_temporaria()
    dados = {
        ges_modelo.GESTOR_SHEET: {c: "" for c in ges_modelo.load_model()[ges_modelo.GESTOR_SHEET]}
    }
    dados[ges_modelo.GESTOR_SHEET]["Código do Inep"] = "22222222"
    dados[ges_modelo.GESTOR_SHEET]["Unidade Escolar"] = "Escola Dois"
    ges_exporter.append_school_temp(dados, pasta)

    estado = _EstadoFalso(
        cg_stop=None,
        cg_lock=threading.Lock(),
        cg_shared={"error": None, "failures": FALHAS},
        cg_tmp_dir=pasta,
        cg_running=True,
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(ges_frontend, "st", st_falso)

    ges_frontend._finalize(1)

    assert estado.cg_running is False
    assert estado.cg_finalized is True

    xlsx, nome = estado.cg_xlsx
    assert nome.endswith(".xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(xlsx))[ges_modelo.GESTOR_SHEET]
    assert "22222222" in [c.value for row in ws.iter_rows() for c in row]
    assert not pasta.exists()

    pend, pend_nome = estado.cg_pending
    assert pend_nome.startswith("vinculo_gestor_escolas_nao_coletadas_")
    ws_pend = openpyxl.load_workbook(io.BytesIO(pend)).active
    assert [c.value for c in ws_pend[2]] == ["222", "Escola Dois", "Falha no login"]

    assert [d["label"] for d in st_falso.downloads] == [
        "Baixar XLSX",
        "Baixar escolas não coletadas (XLSX)",
    ]


def test_finalize_gestor_sem_coleta_limpa_temp_e_mantem_pendentes(monkeypatch):
    pasta = ges_exporter.nova_pasta_temporaria()

    estado = _EstadoFalso(
        cg_stop=None,
        cg_lock=threading.Lock(),
        cg_shared={"error": None, "failures": FALHAS},
        cg_tmp_dir=pasta,
        cg_running=True,
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(ges_frontend, "st", st_falso)

    ges_frontend._finalize(0)

    assert estado.cg_xlsx is None
    assert not pasta.exists()
    assert estado.cg_finalized is True and estado.cg_pending is not None
    assert any("Arquivo XLSX não gerado" in m for m in st_falso.markdowns)
    assert [d["label"] for d in st_falso.downloads] == ["Baixar escolas não coletadas (XLSX)"]


def test_finalize_relatorios_gera_zip_e_pendentes(monkeypatch):
    pasta = rel_zipper.nova_pasta_temporaria()
    rel_zipper.append_csv("1", "Escola Um", b"a;b\n1;2", pasta)
    config = {"state_prefix": "rel_", "zip_prefix": "relatorios_turmas"}

    estado = _EstadoFalso(
        **{
            "rel_stop": None,
            "rel_lock": threading.Lock(),
            "rel_shared": {"error": None, "failures": FALHAS},
            "rel_tmp_dir": pasta,
            "rel_running": True,
        }
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(rel_frontend, "st", st_falso)

    rel_frontend._finalize(config, 1)

    assert estado["rel_running"] is False
    assert estado["rel_finalized"] is True

    zip_bytes, nome = estado["rel_zip"]
    assert zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist() == ["1_Escola Um.csv"]
    assert not pasta.exists()

    pend, pend_nome = estado["rel_pending"]
    assert pend_nome.startswith("relatorios_turmas_escolas_nao_coletadas_")
    ws_pend = openpyxl.load_workbook(io.BytesIO(pend)).active
    assert [c.value for c in ws_pend[2]] == ["222", "Escola Dois", "Falha no login"]

    assert [d["label"] for d in st_falso.downloads] == [
        "Baixar ZIP",
        "Baixar escolas não coletadas (XLSX)",
    ]


def test_finalize_relatorios_sem_coleta_limpa_temp_e_mantem_pendentes(monkeypatch):
    pasta = rel_zipper.nova_pasta_temporaria()
    config = {"state_prefix": "rel_", "zip_prefix": "relatorios_turmas"}

    estado = _EstadoFalso(
        **{
            "rel_stop": None,
            "rel_lock": threading.Lock(),
            "rel_shared": {"error": None, "failures": FALHAS},
            "rel_tmp_dir": pasta,
            "rel_running": True,
        }
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(rel_frontend, "st", st_falso)

    rel_frontend._finalize(config, 0)

    assert estado["rel_zip"] is None
    assert not pasta.exists()
    assert estado["rel_finalized"] is True and estado["rel_pending"] is not None
    assert any("Arquivo Zip não gerado" in m for m in st_falso.markdowns)
    assert [d["label"] for d in st_falso.downloads] == ["Baixar escolas não coletadas (XLSX)"]


def test_finalize_recibos_gera_zip_e_pendentes(monkeypatch):
    pasta = rec_zipper.nova_pasta_temporaria()
    rec_zipper.append_pdf("1", "Escola Um", b"%PDF-1.4\n%%EOF", pasta)
    config = {"state_prefix": "rec_", "zip_prefix": "recibos_fechamento"}

    estado = _EstadoFalso(
        **{
            "rec_stop": None,
            "rec_lock": threading.Lock(),
            "rec_shared": {"error": None, "failures": FALHAS},
            "rec_tmp_dir": pasta,
            "rec_running": True,
        }
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(rec_frontend, "st", st_falso)

    rec_frontend._finalize(config, 1)

    assert estado["rec_running"] is False
    assert estado["rec_finalized"] is True

    zip_bytes, nome = estado["rec_zip"]
    assert zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist() == ["1_Escola Um.pdf"]
    assert not pasta.exists()

    pend, pend_nome = estado["rec_pending"]
    assert pend_nome.startswith("recibos_fechamento_escolas_nao_coletadas_")
    ws_pend = openpyxl.load_workbook(io.BytesIO(pend)).active
    assert [c.value for c in ws_pend[2]] == ["222", "Escola Dois", "Falha no login"]

    assert [d["label"] for d in st_falso.downloads] == [
        "Baixar ZIP",
        "Baixar escolas não coletadas (XLSX)",
    ]


def test_finalize_recibos_sem_coleta_limpa_temp_e_mantem_pendentes(monkeypatch):
    pasta = rec_zipper.nova_pasta_temporaria()
    config = {"state_prefix": "rec_", "zip_prefix": "recibos_fechamento"}

    estado = _EstadoFalso(
        **{
            "rec_stop": None,
            "rec_lock": threading.Lock(),
            "rec_shared": {"error": None, "failures": FALHAS},
            "rec_tmp_dir": pasta,
            "rec_running": True,
        }
    )
    st_falso = _StFalso(estado)
    monkeypatch.setattr(rec_frontend, "st", st_falso)

    rec_frontend._finalize(config, 0)

    assert estado["rec_zip"] is None
    assert not pasta.exists()
    assert estado["rec_finalized"] is True and estado["rec_pending"] is not None
    assert any("Arquivo Zip não gerado" in m for m in st_falso.markdowns)
    assert [d["label"] for d in st_falso.downloads] == ["Baixar escolas não coletadas (XLSX)"]
