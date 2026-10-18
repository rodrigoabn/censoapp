"""Falhas de coleta de gestores (login, escola não encontrada, erro por escola)."""

from __future__ import annotations

import io

import openpyxl
import pytest

from apps.coleta_gestor.src import exporter, gestor_scraper

ESCOLAS = [
    {"codigo_inep": "111", "nome_unidade": "Escola Um"},
    {"codigo_inep": "222", "nome_unidade": "Escola Dois"},
]


class _LocatorFalso:
    def is_visible(self):
        return False


class _PaginaFalsa:
    def locator(self, *a, **k):
        return _LocatorFalso()

    def close(self):
        pass


class _ContextoFalso:
    def new_page(self):
        return _PaginaFalsa()

    def clear_cookies(self):
        pass

    def close(self):
        pass


class _BrowserFalso:
    def new_context(self, **kwargs):
        return _ContextoFalso()


@pytest.fixture(autouse=True)
def _sem_espera(monkeypatch):
    monkeypatch.setattr("time.sleep", lambda *_: None)


def test_login_falho_marca_todas_as_escolas(monkeypatch):
    def _falha(*a, **k):
        raise gestor_scraper.LoginFalhou()

    monkeypatch.setattr(gestor_scraper, "_login_and_read", _falha)
    outputs, failures = gestor_scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS)
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Falha no login"},
        {"codigo_inep": "222", "nome_unidade": "Escola Dois", "motivo": "Falha no login"},
    ]


def test_escola_nao_encontrada_vira_falha(monkeypatch):
    monkeypatch.setattr(gestor_scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(gestor_scraper, "_session_expirada", lambda page: False)
    monkeypatch.setattr(gestor_scraper, "scrape_gestor_school", lambda *a, **k: None)
    outputs, failures = gestor_scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS[:1])
    assert outputs == []
    assert failures == [
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Escola não encontrada"}
    ]


def test_escola_ok_e_erro_inesperado_convivem(monkeypatch):
    monkeypatch.setattr(gestor_scraper, "_login_and_read", lambda *a, **k: None)
    monkeypatch.setattr(gestor_scraper, "_session_expirada", lambda page: False)

    def _scrape(page, codigo, nome):
        if codigo == "222":
            raise RuntimeError("boom")
        return {"Vínculo do gestor": {"Código do Inep": codigo}}

    monkeypatch.setattr(gestor_scraper, "scrape_gestor_school", _scrape)
    outputs, failures = gestor_scraper.run_scraping(_BrowserFalso(), "l", "s", ESCOLAS)
    assert [e["codigo_inep"] for e, _ in outputs] == ["111"]
    assert failures == [
        {"codigo_inep": "222", "nome_unidade": "Escola Dois", "motivo": "Erro: boom"}
    ]


def test_planilha_de_nao_coletadas_tem_cabecalho_e_linhas():
    data = exporter.build_pending_xlsx([
        {"codigo_inep": "111", "nome_unidade": "Escola Um", "motivo": "Falha no login"},
    ])
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["Código do Inep", "Nome da Unidade", "Motivo"]
    assert [c.value for c in ws[2]] == ["111", "Escola Um", "Falha no login"]
    assert exporter.pending_filename().startswith("vinculo_gestor_escolas_nao_coletadas_")
