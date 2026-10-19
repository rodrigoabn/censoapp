"""Temporários, planilhas de pendências e ZIPs compartilhados (Fase 4.4).

Congrega os helpers 1:1 dos quatro coletores (rel/rec `zipper.py` e cad/ges
`exporter.py`), parametrizando prefixo de pasta, qualificador do nome final e
extensão (`.csv`/`.pdf`/`.xlsx`/`.zip`). Os módulos dos apps mantêm os nomes
históricos como wrappers/aliases.

Nenhuma escrita é feita no site Educacenso (cláusula pétrea de modo leitura).
"""

from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def nova_pasta_temporaria(prefix: str) -> Path:
    """Cria a pasta temporária exclusiva desta execução de coleta."""
    return Path(tempfile.mkdtemp(prefix=prefix))


def remover_temp(tmp_dir: Path) -> None:
    """Remove a pasta temporária da execução (se existir)."""
    shutil.rmtree(tmp_dir, ignore_errors=True)


def _file_date() -> str:
    return datetime.now().strftime("%d-%m-%Y")


def output_filename(prefix: str, *, qualificador: str | None = None, ext: str = ".zip") -> str:
    """Nome final do artefato: prefixo [+ qualificador] + data; extensão configurável."""
    partes = [prefix] + ([qualificador] if qualificador else []) + [_file_date()]
    return "_".join(partes) + ext


def pending_filename(prefix: str) -> str:
    return f"{prefix}_escolas_nao_coletadas_{_file_date()}.xlsx"


def _neutralizar_formulas(ws) -> None:
    """Neutraliza textos iniciados com '=' para não virarem fórmula no Excel."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = "'" + cell.value


def _ajustar_larguras(ws, min_width: int = 10, max_width: int = 50, pad: int = 2) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        comps = [len(str(c.value)) for c in col if c.value is not None]
        if not comps:
            ws.column_dimensions[letter].width = min_width
            continue
        ws.column_dimensions[letter].width = min(max(max(comps) + pad, min_width), max_width)


def build_pending_xlsx(failures: list[dict]) -> bytes:
    """Gera um XLSX com as escolas processadas que não tiveram dados coletados.

    Colunas: Código do Inep | Nome da Unidade | Motivo.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Escolas não coletadas"
    ws.append(["Código do Inep", "Nome da Unidade", "Motivo"])
    for f in failures:
        ws.append([
            f.get("codigo_inep", ""),
            f.get("nome_unidade", ""),
            f.get("motivo", ""),
        ])

    _neutralizar_formulas(ws)
    _ajustar_larguras(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _safe_name(name: str) -> str:
    """Sanitiza o nome de arquivo (remove caracteres inválidos)."""
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name or "").strip()
    return name[:120] if name else "escola"


def arquivo_filename(codigo: str, nome: str, ext: str) -> str:
    codigo = _safe_name(codigo)
    nome = _safe_name(nome)
    return f"{codigo}_{nome}{ext}" if nome else f"{codigo}{ext}"


def append_arquivo(codigo: str, nome: str, data: bytes, tmp_dir: Path, ext: str) -> None:
    """Grava o arquivo de UMA escola na pasta temporária."""
    tmp_dir.mkdir(parents=True, exist_ok=True)
    path = tmp_dir / arquivo_filename(codigo, nome, ext)
    # Evita sobrescrever caso haja código repetido.
    if path.exists():
        stem, suffix = path.stem, path.suffix
        i = 2
        while (tmp_dir / f"{stem}_{i}{suffix}").exists():
            i += 1
        path = tmp_dir / f"{stem}_{i}{suffix}"
    path.write_bytes(data)


def count_temp(tmp_dir: Path, ext: str) -> int:
    if not tmp_dir.exists():
        return 0
    return sum(1 for _ in tmp_dir.glob(f"*{ext}"))


def build_zip(tmp_dir: Path, ext: str) -> bytes:
    """Compacta todos os temporários em um ZIP e retorna os bytes.

    Remove os temporários ao final.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        if tmp_dir.exists():
            for path in sorted(tmp_dir.glob(f"*{ext}")):
                zf.write(path, arcname=path.name)
    buffer.seek(0)

    if tmp_dir.exists():
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return buffer.getvalue()
