from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from datetime import date
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


SHEETS = {
    "unidades": "Cadastro da Unidade",
    "gestores": "Gestores",
    "turmas": "Turmas",
    "alunos": "Alunos",
    "professores": "Professores",
}

_IGNORED_MESSAGE = "Relatório não gerado"
_NO_ISSUES_MESSAGE = "Unidade não possui inconsistência"


def _text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value))
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return "".join(char.lower() for char in normalized if char.isalnum())


def _normalize_inep(value: Any, location: str) -> str:
    code = _text(value)
    if re.fullmatch(r"33\d{6}\.0+", code):
        code = code.split(".", 1)[0]
    if not re.fullmatch(r"33\d{6}", code):
        raise ValueError(
            f"Código INEP inválido em {location}: {code!r}. "
            "O código deve ter 8 dígitos e iniciar com 33."
        )
    return code


def read_active_schools(raw_bytes: bytes) -> list[dict[str, str]]:
    try:
        frame = pd.read_excel(io.BytesIO(raw_bytes), dtype=object)
    except Exception as exc:
        raise ValueError(f"Não foi possível ler a relação de escolas ativas: {exc}") from exc

    if frame.shape[1] < 2:
        raise ValueError(
            "A relação de escolas ativas deve conter pelo menos duas colunas: "
            "código INEP na primeira e nome da unidade na segunda."
        )

    schools: list[dict[str, str]] = []
    seen_codes: set[str] = set()
    for row_index, row in frame.iterrows():
        raw_code = row.iloc[0]
        raw_name = row.iloc[1]
        if (raw_code is None or pd.isna(raw_code)) and (
            raw_name is None or pd.isna(raw_name)
        ):
            continue

        location = f"linha {row_index + 2} da relação de escolas ativas"
        code = _normalize_inep(raw_code, location)
        name = _text(raw_name)
        if not name:
            raise ValueError(f"Nome da unidade vazio em {location}.")
        if code in seen_codes:
            raise ValueError(f"Código INEP duplicado na relação de escolas ativas: {code}.")
        seen_codes.add(code)
        schools.append({"codigo_inep": code, "nome_unidade": name})

    if not schools:
        raise ValueError("A relação de escolas ativas não contém unidades válidas.")
    return schools


def _read_report(raw_bytes: bytes, label: str) -> tuple[list[str], dict[str, list[dict[str, Any]]]]:
    try:
        sheets = pd.read_excel(io.BytesIO(raw_bytes), sheet_name=None, dtype=object)
    except Exception as exc:
        raise ValueError(f"Não foi possível ler o relatório '{label}': {exc}") from exc

    columns: list[str] = []
    rows_by_code: dict[str, list[dict[str, Any]]] = {}
    found_code_column = False

    for sheet_name, frame in sheets.items():
        if not len(frame.columns):
            continue
        code_column_index = next(
            (
                index
                for index, column in enumerate(frame.columns)
                if "codigo" in _normalize_header(column)
                and "inep" in _normalize_header(column)
            ),
            None,
        )
        if code_column_index is None:
            raise ValueError(
                f"O relatório '{label}', aba '{sheet_name}', não possui uma coluna "
                "de código INEP (por exemplo, 'Código do Inep')."
            )
        found_code_column = True

        for column in frame.columns:
            name = str(column)
            if name not in columns:
                columns.append(name)

        for row_index, row in frame.iterrows():
            values = row.tolist()
            if not any(_text(value) for value in values):
                continue
            location = f"linha {row_index + 2} da aba '{sheet_name}' do relatório '{label}'"
            code = _normalize_inep(values[code_column_index], location)
            record = {
                str(column): value
                for column, value in zip(frame.columns, values)
            }
            rows_by_code.setdefault(code, []).append(record)

    if not found_code_column:
        raise ValueError(
            f"O relatório '{label}' não contém linhas com cabeçalho e código INEP."
        )
    return columns, rows_by_code


def _excel_value(value: Any) -> Any:
    if value is None or pd.isna(value):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


def _write_report_sheet(
    workbook: openpyxl.Workbook,
    title: str,
    columns: list[str],
    rows: list[dict[str, Any]],
    message: str | None,
) -> None:
    worksheet = workbook.create_sheet(title)
    if message:
        worksheet.cell(row=1, column=1, value=message)
        worksheet.column_dimensions["A"].width = min(max(len(message) + 2, 20), 60)
        return

    worksheet.append([_excel_value(column) for column in columns])
    for row in rows:
        worksheet.append([_excel_value(row.get(column)) for column in columns])

    worksheet.freeze_panes = "A2"
    if columns:
        worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, column in enumerate(worksheet.columns, start=1):
        values = [len(str(cell.value)) for cell in column if cell.value is not None]
        width = min(max(max(values, default=10) + 2, 10), 50)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def build_school_reports_zip(
    active_schools_bytes: bytes,
    report_bytes: dict[str, bytes | None],
    ignored_reports: set[str],
) -> tuple[bytes, dict[str, Any]]:
    """Build one XLSX per active school and return aggregate statistics."""
    unknown_reports = set(report_bytes) - set(SHEETS)
    unknown_ignored = ignored_reports - set(SHEETS)
    if unknown_reports or unknown_ignored:
        raise ValueError("Foram informadas categorias de relatório desconhecidas.")

    report_data: dict[str, tuple[list[str], dict[str, list[dict[str, Any]]]]] = {}
    for key, title in SHEETS.items():
        if key in ignored_reports:
            continue
        raw_bytes = report_bytes.get(key)
        if raw_bytes is None:
            raise ValueError(
                f"Envie o relatório '{title}' ou marque a opção "
                "'Não usar este relatório'."
            )
        report_data[key] = _read_report(raw_bytes, title)

    schools = read_active_schools(active_schools_bytes)
    filenames: set[str] = set()
    school_filenames: list[tuple[dict[str, str], str, bool]] = []
    stats: dict[str, Any] = {
        "total_units": len(schools),
        "units_with_inconsistencies": 0,
        "units_without_inconsistencies": 0,
        "by_type": {
            key: {
                "label": title,
                "used": key in report_data,
                "units": 0,
                "issues": 0,
            }
            for key, title in SHEETS.items()
        },
        "per_school": [],
    }

    for school in schools:
        per_type_counts = {}
        for key in SHEETS:
            if key in report_data:
                _, records_by_code = report_data[key]
                per_type_counts[key] = len(
                    records_by_code.get(school["codigo_inep"], [])
                )
            else:
                per_type_counts[key] = 0
        total_issues = sum(per_type_counts.values())
        has_inconsistencies = total_issues > 0
        if has_inconsistencies:
            stats["units_with_inconsistencies"] += 1
        else:
            stats["units_without_inconsistencies"] += 1

        for key, count in per_type_counts.items():
            stats["by_type"][key]["issues"] += count
            if count:
                stats["by_type"][key]["units"] += 1

        stats["per_school"].append({
            "codigo_inep": school["codigo_inep"],
            "nome_unidade": school["nome_unidade"],
            "total_issues": total_issues,
            "by_type": per_type_counts,
        })

        filename_suffix = "" if has_inconsistencies else "(SEM INCONSISTENCIAS)"
        filename = f"{_safe_filename(school['nome_unidade'])}{filename_suffix}.xlsx"
        normalized_filename = filename.casefold()
        if normalized_filename in filenames:
            raise ValueError(
                "Há nomes de unidades escolares repetidos ou iguais após "
                "a normalização do nome do arquivo. Corrija a relação de escolas."
            )
        filenames.add(normalized_filename)
        school_filenames.append((school, filename, has_inconsistencies))

    archive_buffer = io.BytesIO()
    with zipfile.ZipFile(archive_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for school, filename, has_inconsistencies in school_filenames:
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)

            for key, title in SHEETS.items():
                if key in ignored_reports:
                    _write_report_sheet(workbook, title, [], [], _IGNORED_MESSAGE)
                    continue

                columns, records_by_code = report_data[key]
                rows = records_by_code.get(school["codigo_inep"], [])
                if has_inconsistencies and not rows:
                    continue
                message = None if rows else _NO_ISSUES_MESSAGE
                _write_report_sheet(workbook, title, columns, rows, message)

            workbook_buffer = io.BytesIO()
            workbook.save(workbook_buffer)
            archive.writestr(filename, workbook_buffer.getvalue())
    return archive_buffer.getvalue(), stats


def _safe_filename(name: str) -> str:
    normalized = unicodedata.normalize("NFC", name)
    safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", normalized).strip().rstrip(".")
    return safe_name[:180] or "unidade_escolar"


def output_filename() -> str:
    return f"unificador_inconsistencias_{date.today():%d-%m-%Y}.zip"
