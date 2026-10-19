"""Exportação local dos dados coletados (apenas-leitura no site).

Gera um único arquivo .xlsx com a aba "Vínculo do gestor" (espelho de
modelo_coleta_gestor.xlsx). Durante a raspagem, cada escola é appendada em
um CSV temporário; ao final, o CSV é consolidado no XLSX final.

Nenhuma escrita é feita no site Educacenso (cláusula pétrea de modo leitura).
"""

from __future__ import annotations

import csv
import io
import shutil
from pathlib import Path

import openpyxl

from apps.common import export as _export

from . import modelo

_TMP_PREFIX = "educacenso_gestor_tmp_"


def nova_pasta_temporaria() -> Path:
    """Wrapper de `apps.common.export` (histórico): pasta exclusiva da execução."""
    return _export.nova_pasta_temporaria(_TMP_PREFIX)


remover_temp = _export.remover_temp


SHEET = modelo.GESTOR_SHEET


_file_date = _export._file_date


def output_filename() -> str:
    """Wrapper de `apps.common.export` (histórico): nome do XLSX final do app."""
    return _export.output_filename("vinculo_gestor", qualificador="educacenso", ext=".xlsx")


def pending_filename() -> str:
    """Wrapper de `apps.common.export` (histórico): nome da planilha de não coletadas."""
    return _export.pending_filename("vinculo_gestor")


_neutralizar_formulas = _export._neutralizar_formulas
_ajustar_larguras = _export._ajustar_larguras
build_pending_xlsx = _export.build_pending_xlsx


def _sheet_file(sheet: str) -> str:
    """Nome de arquivo seguro para a aba (mantém legível, sem caracteres inválidos)."""
    name = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in sheet)
    return f"{name}.csv"


def append_school_temp(dados: dict[str, dict], tmp_dir: Path) -> None:
    """Append os dados de UMA escola em um CSV temporário por aba do modelo.

    `dados` é {nome_da_aba: {coluna: valor}}.
    """
    tmp_dir.mkdir(parents=True, exist_ok=True)
    model = modelo.load_model()
    for sheet, cols in model.items():
        if not cols:
            continue
        row = dados.get(sheet, {c: "" for c in cols})
        path = tmp_dir / _sheet_file(sheet)
        new = not path.exists()
        with path.open("a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=cols)
            if new:
                writer.writeheader()
            writer.writerow({c: row.get(c, "") for c in cols})


def build_xlsx(tmp_dir: Path, model=None) -> bytes:
    """Consolida os CSVs temporários em um único XLSX (todas as abas do modelo)
    e retorna os bytes. Remove os CSVs temporários ao final.
    """
    if model is None:
        model = modelo.load_model()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba em branco padrão

    for sheet, cols in model.items():
        if not cols:
            continue
        ws = wb.create_sheet(title=sheet)
        ws.append(cols)  # cabeçalho = 1ª linha do modelo
        path = tmp_dir / _sheet_file(sheet)
        if path.exists():
            with path.open(encoding="utf-8") as f:
                reader = csv.reader(f)
                next(reader, None)  # pula cabeçalho do CSV
                for r in reader:
                    ws.append(r)

    for sheet in wb.sheetnames:
        _ajustar_larguras(wb[sheet])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    if tmp_dir.exists():
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return buffer.getvalue()
