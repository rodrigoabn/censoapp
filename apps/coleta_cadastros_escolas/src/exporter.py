"""Exportação local dos dados coletados (apenas-leitura no site).

Gera um único arquivo .xlsx com 6 abas (espelho de modelo_coleta.xlsx).
Durante a raspagem, cada escola é appendada em CSVs temporários por aba
(resiliência). Ao final, os CSVs são consolidados no XLSX final.

Nenhuma escrita é feita no site Educacenso (cláusula pétrea de modo leitura).
"""

from __future__ import annotations

import csv
import io
import shutil
from pathlib import Path

import openpyxl

from apps.common import export as _export

from . import modelo

_TMP_PREFIX = "educacenso_tmp_"


def nova_pasta_temporaria() -> Path:
    """Wrapper de `apps.common.export` (histórico): pasta exclusiva da execução."""
    return _export.nova_pasta_temporaria(_TMP_PREFIX)


remover_temp = _export.remover_temp
_file_date = _export._file_date


def output_filename(prefix: str | None = None) -> str:
    """Wrapper de `apps.common.export` (histórico): nome do XLSX final do app.

    O prefixo opcional atende à chamada genérica de `apps.common.ui`; sem ele,
    vale o prefixo histórico do app.
    """
    return _export.output_filename(
        prefix or "cadastro_escola", qualificador="educacenso", ext=".xlsx"
    )


def pending_filename(prefix: str | None = None) -> str:
    """Wrapper de `apps.common.export` (histórico): nome da planilha de não coletadas."""
    return _export.pending_filename(prefix or "cadastro_escola")


_neutralizar_formulas = _export._neutralizar_formulas
_ajustar_larguras = _export._ajustar_larguras
build_pending_xlsx = _export.build_pending_xlsx


def _safe_sheet_title(model_title: str) -> str:
    """Excel limita abas a 31 chars e proíbe []:*?/\\."""
    return model_title[:31].replace("/", "-").replace(":", "-").replace("[", "(").replace("]", ")")


def append_school_temp(dados: dict[str, dict], tmp_dir: Path) -> None:
    """Append os dados de UMA escola nos CSVs temporários (um por aba)."""
    tmp_dir.mkdir(parents=True, exist_ok=True)
    model = modelo.load_model()
    for sheet, row in dados.items():
        cols = model.get(sheet, [])
        if not cols:
            continue
        path = tmp_dir / f"{_safe_sheet_title(sheet)}.csv"
        new = not path.exists()
        with path.open("a", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=cols)
            if new:
                writer.writeheader()
            writer.writerow({c: row.get(c, "") for c in cols})


def build_xlsx(tmp_dir: Path, model=None) -> bytes:
    """Consolida os CSVs temporários em um único XLSX (6 abas) e retorna bytes.

    Remove os CSVs temporários ao final. O nome do arquivo final é tratado
    pelo chamador (frontend) via output_filename().
    """
    if model is None:
        model = modelo.load_model()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove aba em branco padrão

    for sheet, cols in model.items():
        title = _safe_sheet_title(sheet)
        path = tmp_dir / f"{title}.csv"
        ws = wb.create_sheet(title=title)
        ws.append(cols)  # cabeçalho = 2ª linha do modelo
        if path.exists():
            with path.open(encoding="utf-8") as f:
                reader = csv.reader(f)
                next(reader, None)  # pula cabeçalho do CSV
                for r in reader:
                    ws.append(r)

    for sheet in wb.sheetnames:
        _ajustar_larguras(wb[sheet])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Limpeza dos temporários.
    if tmp_dir.exists():
        shutil.rmtree(tmp_dir, ignore_errors=True)

    return buffer.getvalue()
