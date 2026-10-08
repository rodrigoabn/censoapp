from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import streamlit as st

from .checar_estrutura import checar_estrutura
from .upload import (
    checar_colunas,
    read_aluno_zip,
    read_profissionais_zip,
    read_school_list,
    read_turmas_zip,
    read_zip_merged,
)

_PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent.parent
_REFERENCIA_PATHS = {
    "unidades": (
        _PROJECT_ROOT
        / "arquivos de teste"
        / "cadastro_escola_educacenso_26-07-2026.xlsx"
    ),
    "gestor": (
        _PROJECT_ROOT
        / "arquivos de teste"
        / "vinculo_gestor_educacenso_26-07-2026.xlsx"
    ),
}
_REFERENCIA_BYTES = {}
for _key, _path in _REFERENCIA_PATHS.items():
    try:
        _REFERENCIA_BYTES[_key] = _path.read_bytes()
    except Exception:
        _REFERENCIA_BYTES[_key] = None


def _output_filename(config: dict, suffix: str = "inconsistencias", ext: str = "xlsx") -> str:
    data_str = datetime.now().strftime("%d-%m-%Y")
    return f"{suffix}_{config['key']}_{data_str}.{ext}"


def _ajustar_larguras(ws, min_width: int = 10, max_width: int = 50, pad: int = 2) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        compr = [len(str(c.value)) for c in col if c.value is not None]
        if not compr:
            ws.column_dimensions[letter].width = min_width
            continue
        ws.column_dimensions[letter].width = min(max(max(compr) + pad, min_width), max_width)


@st.cache_data(show_spinner=False, max_entries=12)
def _unificado_xlsx(df) -> bytes:
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False, engine="openpyxl")
    buffer.seek(0)
    wb = load_workbook(buffer)
    for ws in wb.worksheets:
        _ajustar_larguras(ws)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def _render_download_grid(opcoes: list[tuple[str, bytes, str, str]]) -> None:
    for i in range(0, len(opcoes), 4):
        cols = st.columns(4)
        for j, col in enumerate(cols):
            idx = i + j
            if idx >= len(opcoes):
                break
            label, data, file_name, key = opcoes[idx]
            with col:
                st.download_button(
                    label,
                    data=data,
                    file_name=file_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key=key,
                    use_container_width=True,
                )


def _modo_relatorio(config: dict) -> str:
    if config.get("split_tipo_turma"):
        return "turma"
    if config.get("split_profissionais"):
        return "profissional"
    if config.get("split_aluno"):
        return "aluno"
    return "padrao"


@st.cache_data(show_spinner=False, max_entries=12)
def _build_xlsx(resultados: list[dict], modo: str = "padrao") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inconsistencias"
    if modo == "profissional":
        max_erros = max((len(r.get("erros") or []) for r in resultados), default=0)
        ws.append(
            ["Codigo do Inep", "Unidade Escolar", "Nome do Profissional", "Identificação Única"]
            + [f"INCONSISTENCIA {i}" for i in range(1, max_erros + 1)]
        )
        for r in resultados:
            erros = list(r.get("erros") or [])
            ws.append([
                r.get("codigo_inep", ""),
                r.get("nome_unidade", ""),
                r.get("nome_profissional", ""),
                r.get("identificacao_unica", ""),
                *erros,
                *([""] * (max_erros - len(erros))),
            ])
    elif modo == "turma":
        max_erros = max((len(r.get("erros") or []) for r in resultados), default=0)
        ws.append(
            ["Codigo do Inep", "Unidade Escolar", "Nome da Turma", "Etapa de ensino",
             "Carga horária semanal (hh:mm)", "periodo_turma", "verifica_integral"]
            + [f"INCONSISTENCIA {i}" for i in range(1, max_erros + 1)]
        )
        for r in resultados:
            erros = list(r.get("erros") or [])
            ws.append([
                r.get("codigo_inep", ""),
                r.get("nome_unidade", ""),
                r.get("nome_turma", ""),
                r.get("etapa_ensino", ""),
                r.get("carga_horaria", ""),
                r.get("periodo_turma", ""),
                r.get("verifica_integral", ""),
                *erros,
                *([""] * (max_erros - len(erros))),
            ])
    elif modo == "aluno":
        max_erros = max((len(r.get("erros") or []) for r in resultados), default=0)
        ws.append(
            ["Codigo do Inep", "Unidade Escolar", "Identificação Única", "Nome do Aluno",
             "Data de nascimento", "Idade", "Nome da Turma", "Etapa de ensino"]
            + [f"INCONSISTENCIA {i}" for i in range(1, max_erros + 1)]
        )
        for r in resultados:
            erros = list(r.get("erros") or [])
            ws.append([
                r.get("codigo_inep", ""),
                r.get("nome_unidade", ""),
                r.get("identificacao_unica", ""),
                r.get("nome", ""),
                r.get("data_nascimento", ""),
                r.get("idade", ""),
                r.get("nome_turma", ""),
                r.get("etapa_ensino", ""),
                *erros,
                *([""] * (max_erros - len(erros))),
            ])
    else:
        ws.append(["Codigo do Inep", "Unidade Escolar", "Aba", "Detalhes"])
        for r in resultados:
            ws.append([
                r.get("codigo_inep", ""),
                r.get("nome_unidade", ""),
                r.get("tipo_validacao", ""),
                r.get("detalhes", ""),
            ])
    _ajustar_larguras(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@st.cache_data(show_spinner=False, max_entries=12)
def _duplicadas_xlsx(registros: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matriculas Duplicadas"
    ws.append([
        "Código do Inep", "Unidade Escolar", "Identificação Única", "Nome do Aluno",
        "CPF", "Motivo", "Código da Matrícula", "Código da Turma", "Nome da Turma",
    ])
    for r in registros:
        ws.append([
            r.get("codigo_inep", ""),
            r.get("nome_unidade", ""),
            r.get("identificacao_unica", ""),
            r.get("nome", ""),
            r.get("cpf", ""),
            r.get("motivo", ""),
            r.get("codigo_matricula", ""),
            r.get("codigo_turma", ""),
            r.get("nome_turma", ""),
        ])
    _ajustar_larguras(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@st.cache_data(show_spinner=False, max_entries=12)
def _levantamento_cor_raca_xlsx(registros: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Levantamento_CorRaca"
    ws.append([
        "Código do Inep", "Nome da Unidade", "Preta", "Parda", "Branca",
        "Amarela", "Indígena", "Não declarado", "Sem informação",
    ])
    for r in registros:
        ws.append([
            r.get("codigo_inep", ""),
            r.get("nome_unidade", ""),
            r.get("preta", 0),
            r.get("parda", 0),
            r.get("branca", 0),
            r.get("amarela", 0),
            r.get("indigena", 0),
            r.get("nao_declarado", 0),
            r.get("sem_informacao", 0),
        ])
    _ajustar_larguras(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@st.cache_data(show_spinner=False, max_entries=12)
def _sem_cpf_xlsx(registros: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alunos_Sem_CPF"
    ws.append([
        "Código do Inep", "Unidade Escolar", "Identificação Única", "Nome do Aluno",
        "CPF",
    ])
    for r in registros:
        ws.append([
            r.get("codigo_inep", ""),
            r.get("nome_unidade", ""),
            r.get("identificacao_unica", ""),
            r.get("nome", ""),
            r.get("cpf", ""),
        ])
    _ajustar_larguras(ws)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _init_state(config: dict) -> None:
    prefix = config["state_prefix"]
    ss = st.session_state
    ss.setdefault(f"{prefix}resultados", None)
    ss.setdefault(f"{prefix}finalizado", False)
    ss.setdefault(f"{prefix}estrutura_ok", False)
    ss.setdefault(f"{prefix}escolas", None)
    ss.setdefault(f"{prefix}ano_letivo", datetime.now().year)
    if config.get("requires_dates", True):
        ss.setdefault(f"{prefix}inicio_ano_letivo", None)
        ss.setdefault(f"{prefix}termino_ano_letivo", None)
    if config.get("source") == "zip":
        ss.setdefault(f"{prefix}df_unificado", None)
    if config.get("split_tipo_turma"):
        ss.setdefault(f"{prefix}df_aee", None)
        ss.setdefault(f"{prefix}df_atv", None)
        ss.setdefault(f"{prefix}df_curricular", None)
        ss.setdefault(f"{prefix}df_escolarizacao", None)
    if config.get("split_profissionais"):
        ss.setdefault(f"{prefix}df_prof_aee_atv", None)
        ss.setdefault(f"{prefix}df_prof_curricular", None)
    if config.get("split_aluno"):
        ss.setdefault(f"{prefix}df_aluno_curricular", None)
        ss.setdefault(f"{prefix}df_aluno_aee", None)
        ss.setdefault(f"{prefix}df_aluno_atv", None)
        ss.setdefault(f"{prefix}df_aluno_transporte", None)
        ss.setdefault(f"{prefix}df_aluno_deficiencia", None)
    ss.setdefault(f"{prefix}colunas_validadas", False)
    ss.setdefault(f"{prefix}colunas_erros", None)
    arquivo_key = f"{prefix}arquivo_bytes"
    if arquivo_key not in ss:
        ss[arquivo_key] = None


def render_validation_frontend(config: dict, logo_base64: str | None = None) -> None:
    _init_state(config)
    prefix = config["state_prefix"]
    is_turmas = config.get("split_tipo_turma", False)
    is_prof = config.get("split_profissionais", False)
    is_aluno = config.get("split_aluno", False)
    is_zip = config.get("source") == "zip"

    if logo_base64:
        banner_html = f"""
        <div class="title-container" style="display:flex; align-items:center; gap:1.5rem;">
            <img src="data:image/png;base64,{logo_base64}"
                 style="width:68px; height:68px; border-radius:12px;
                        box-shadow:0 4px 12px rgba(0,0,0,0.2);
                        background:#fff; padding:5px;" />
            <div>
                <h1 style="margin:0; padding:0; line-height:1.2; color:white !important;">
                    {config['title']}
                </h1>
                <p style="margin-top:0.3rem; color:#bfdbfe; font-size:1rem; margin-bottom:0;">
                    {config['subtitle']}
                </p>
            </div>
        </div>"""
    else:
        banner_html = f"""
        <div class="title-container">
            <h1>{config['title']}</h1>
            <p>{config['subtitle']}</p>
        </div>"""
    st.markdown(banner_html, unsafe_allow_html=True)

    if is_aluno:
        st.markdown("""
        <style>
        ul[data-baseweb="menu"][role="listbox"] {
            width: 5cm !important;
            max-height: 360px !important;
        }
        </style>
        """, unsafe_allow_html=True)

    col_form, col_help = st.columns([1.4, 1])

    with col_form:
        st.markdown('<div class="custom-card">', unsafe_allow_html=True)
        with st.form(f"validacao_{config['key']}"):
            if config.get("requires_dates", True):
                col_data1, col_data2 = st.columns(2)
                with col_data1:
                    inicio = st.date_input(
                        "Inicio do Ano Letivo",
                        value=None,
                        format="DD/MM/YYYY",
                    )
                with col_data2:
                    termino = st.date_input(
                        "Termino do Ano Letivo",
                        value=None,
                        format="DD/MM/YYYY",
                    )
            else:
                inicio = None
                termino = None
            if is_aluno:
                ano_letivo = st.selectbox(
                    "Ano letivo",
                    options=list(range(datetime.now().year, 1999, -1)),
                    index=max(datetime.now().year - st.session_state.get(f"{prefix}ano_letivo", datetime.now().year), 0),
                )
            else:
                ano_letivo = None
            if is_zip:
                arquivo = st.file_uploader(
                    "Arquivo de dados (.zip)",
                    type=["zip"],
                    help="ZIP contendo os CSVs exportados do Educacenso (um por escola).",
                )
            else:
                arquivo = st.file_uploader(
                    "Arquivo de dados (.xlsx / .xls)",
                    type=["xlsx", "xls"],
                    help="Deve conter as colunas: 'codigo do inep' e 'nome da unidade'.",
                )
            submitted = st.form_submit_button("Carregar e processar" if is_zip else "Verificar estrutura")
        st.markdown('</div>', unsafe_allow_html=True)

    with col_help:
        st.markdown('<div class="custom-card">', unsafe_allow_html=True)
        st.markdown("""
        <h3>Orientacoes de uso</h3>
        """, unsafe_allow_html=True)
        if is_zip and is_aluno:
            st.markdown("""
        1. Envie o arquivo ZIP com os CSVs exportados do Educacenso (um por escola).
        2. Clique em **Carregar e processar**: unifica e valida as colunas automaticamente.
        3. Baixe a **relação completa** e as tabelas de **alunos** (curricular, AEE, Atv Complementar, Transporte, Deficiência).
        """, unsafe_allow_html=True)
        elif is_zip:
            st.markdown("""
        1. Envie o arquivo ZIP com os CSVs exportados do Educacenso (um por escola).
        2. Clique em **Carregar e processar**: unifica e valida as colunas automaticamente.
        3. Baixe a **relação completa** e, em Turmas, as tabelas de **AEE**, **Atividade Complementar** e **Escolarização** (XLSX/CSV).
        4. Clique em **Iniciar Verificação de Inconsistências**.
        5. O relatorio de inconsistencias fica disponivel para download.
        """, unsafe_allow_html=True)
        else:
            st.markdown("""
        1. Envie o arquivo XLSX exportado do Educacenso.
        2. Clique em **Verificar estrutura** para validar abas e colunas.
        3. Apos aprovado, clique em **Iniciar validacoes**.
        4. O relatorio de inconsistencias fica disponivel para download.
        """, unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    estrutura_ok = st.session_state.get(f"{prefix}estrutura_ok", False)
    escolas = st.session_state.get(f"{prefix}escolas")
    resultados = st.session_state.get(f"{prefix}resultados")
    finalizado = st.session_state.get(f"{prefix}finalizado", False)

    # Etapa 2: estrutura ok -> resumo + validacao de colunas + downloads + botao inconsistências
    if estrutura_ok and escolas is not None:
        if is_zip:
            df_unificado = st.session_state.get(f"{prefix}df_unificado")
            n_unidades = len(escolas)
            n_registros = len(df_unificado) if df_unificado is not None else 0
            if is_turmas:
                qtd_aee = len(st.session_state.get(f"{prefix}df_aee")) if st.session_state.get(f"{prefix}df_aee") is not None else 0
                qtd_atv = len(st.session_state.get(f"{prefix}df_atv")) if st.session_state.get(f"{prefix}df_atv") is not None else 0
                qtd_curricular = len(st.session_state.get(f"{prefix}df_curricular")) if st.session_state.get(f"{prefix}df_curricular") is not None else 0
                qtd_escol = len(st.session_state.get(f"{prefix}df_escolarizacao")) if st.session_state.get(f"{prefix}df_escolarizacao") is not None else 0
                st.success(
                    f"ZIP processado com sucesso. {n_unidades} unidade(s), "
                    f"{n_registros} registro(s) totais. Tabelas geradas: "
                    f"{qtd_aee} em turmas_aee, {qtd_atv} em turmas_atvcomplementar, "
                    f"{qtd_curricular} em turmas_curricular, "
                    f"{qtd_escol} em turmas_escolarizacao."
                )
            elif is_prof:
                qtd_aee = len(st.session_state.get(f"{prefix}df_prof_aee_atv")) if st.session_state.get(f"{prefix}df_prof_aee_atv") is not None else 0
                qtd_cur = len(st.session_state.get(f"{prefix}df_prof_curricular")) if st.session_state.get(f"{prefix}df_prof_curricular") is not None else 0
                st.success(
                    f"ZIP processado com sucesso. {n_unidades} unidade(s), "
                    f"{n_registros} registro(s) totais. Tabelas geradas: "
                    f"{qtd_aee} em profissionais_aee_atv, "
                    f"{qtd_cur} em profissionais_curricular."
                )
            elif is_aluno:
                qtd_pcur = len(st.session_state.get(f"{prefix}df_aluno_curricular")) if st.session_state.get(f"{prefix}df_aluno_curricular") is not None else 0
                qtd_paee = len(st.session_state.get(f"{prefix}df_aluno_aee")) if st.session_state.get(f"{prefix}df_aluno_aee") is not None else 0
                qtd_patv = len(st.session_state.get(f"{prefix}df_aluno_atv")) if st.session_state.get(f"{prefix}df_aluno_atv") is not None else 0
                qtd_ptransp = len(st.session_state.get(f"{prefix}df_aluno_transporte")) if st.session_state.get(f"{prefix}df_aluno_transporte") is not None else 0
                qtd_pdef = len(st.session_state.get(f"{prefix}df_aluno_deficiencia")) if st.session_state.get(f"{prefix}df_aluno_deficiencia") is not None else 0
                st.success(
                    f"ZIP processado com sucesso. {n_unidades} unidade(s), "
                    f"{n_registros} registro(s) totais. Tabelas geradas: "
                    f"{qtd_pcur} em aluno_curricular, {qtd_paee} em aluno_aee, "
                    f"{qtd_patv} em aluno_atv_complementar, "
                    f"{qtd_ptransp} em aluno_transporte, "
                    f"{qtd_pdef} em aluno_deficiencia."
                )
            else:
                st.success(f"ZIP processado com sucesso. {n_unidades} unidade(s), {n_registros} registro(s).")

            colunas_validadas = st.session_state.get(f"{prefix}colunas_validadas", False)
            colunas_erros = st.session_state.get(f"{prefix}colunas_erros")

            if not colunas_validadas:
                if colunas_erros is not None:
                    faltando, extras = colunas_erros
                    if faltando:
                        st.error(f"Colunas faltantes: {', '.join(faltando)}")
                    if extras:
                        st.error(f"Colunas não esperadas: {', '.join(extras)}")
                st.info("As colunas do arquivo unificado não correspondem ao esperado. "
                        "Reenvie um arquivo corrigido para prosseguir.")

            if colunas_validadas:
                df_aee = st.session_state.get(f"{prefix}df_aee") if is_turmas else None
                df_atv = st.session_state.get(f"{prefix}df_atv") if is_turmas else None
                df_escol = st.session_state.get(f"{prefix}df_escolarizacao") if is_turmas else None

                if is_aluno:
                    listagem_label = "Carregando Listagem de Alunos..."
                elif is_turmas:
                    listagem_label = "Carregando Listagem de Turmas..."
                elif is_prof:
                    listagem_label = "Carregando Listagem de Profissionais..."
                else:
                    listagem_label = "Carregando Listagem de Dados..."

                with st.spinner(listagem_label):
                    opcoes: list[tuple[str, bytes, str, str]] = []

                    if df_unificado is not None:
                        opcoes.append((
                            "Relação completa",
                            _unificado_xlsx(df_unificado),
                            _output_filename(config, suffix=f"{config['key']}_relacao_completa", ext="xlsx"),
                            f"{prefix}btn_xlsx",
                        ))

                    if is_prof:
                        df_paee = st.session_state.get(f"{prefix}df_prof_aee_atv")
                        df_pcur = st.session_state.get(f"{prefix}df_prof_curricular")
                        for label, df, suffix, key in [
                            ("Profissionais AEE / Atv Complementar", df_paee, "prof_aee_atv", f"{prefix}btn_profaee_xlsx"),
                            ("Profissionais Curricular", df_pcur, "prof_curricular", f"{prefix}btn_profcur_xlsx"),
                        ]:
                            if df is not None:
                                opcoes.append((
                                    label,
                                    _unificado_xlsx(df),
                                    _output_filename(config, suffix=suffix, ext="xlsx"),
                                    key,
                                ))

                    if is_aluno:
                        df_pcur = st.session_state.get(f"{prefix}df_aluno_curricular")
                        df_paee = st.session_state.get(f"{prefix}df_aluno_aee")
                        df_patv = st.session_state.get(f"{prefix}df_aluno_atv")
                        df_ptransp = st.session_state.get(f"{prefix}df_aluno_transporte")
                        df_pdef = st.session_state.get(f"{prefix}df_aluno_deficiencia")
                        for label, df, suffix, key in [
                            ("Alunos Curriculares", df_pcur, "aluno_curricular", f"{prefix}btn_aluno_cur_xlsx"),
                            ("Alunos AEE", df_paee, "aluno_aee", f"{prefix}btn_aluno_aee_xlsx"),
                            ("Alunos Atv Complementar", df_patv, "aluno_atv_complementar", f"{prefix}btn_aluno_atv_xlsx"),
                            ("Alunos Transporte Escolar", df_ptransp, "aluno_transporte", f"{prefix}btn_aluno_transp_xlsx"),
                            ("Alunos Deficiência / Altas Habilidades", df_pdef, "aluno_deficiencia", f"{prefix}btn_aluno_def_xlsx"),
                        ]:
                            if df is not None:
                                opcoes.append((
                                    label,
                                    _unificado_xlsx(df),
                                    _output_filename(config, suffix=suffix, ext="xlsx"),
                                    key,
                                ))

                    if is_turmas:
                        for label, df, suffix, key in [
                            ("Turmas AEE", df_aee, "turmas_aee", f"{prefix}btn_aee_xlsx"),
                            ("Turmas Atv Complementar", df_atv, "turmas_atv_complementar", f"{prefix}btn_atv_xlsx"),
                            ("Turmas Escolarização", df_escol, "turmas_escolarizacao", f"{prefix}btn_escol_xlsx"),
                        ]:
                            if df is not None:
                                opcoes.append((
                                    label,
                                    _unificado_xlsx(df),
                                    _output_filename(config, suffix=suffix, ext="xlsx"),
                                    key,
                                ))

                if is_aluno:
                    st.markdown("**LISTAGEM DE ALUNOS**")

                _render_download_grid(opcoes)

                if is_aluno:
                    st.markdown("**RELATÓRIOS DE APOIO**")
                    with st.spinner("Carregando Relatórios de Apoio..."):
                        find_duplicadas = config.get("find_duplicadas")
                        df_aluno_cur = st.session_state.get(f"{prefix}df_aluno_curricular")
                        if find_duplicadas is not None and df_aluno_cur is not None:
                            duplicadas = find_duplicadas(df_aluno_cur)
                            if duplicadas:
                                st.download_button(
                                    "Matrículas em Duplicidade (XLSX)",
                                    data=_duplicadas_xlsx(duplicadas),
                                    file_name=_output_filename(config, suffix="matriculas_duplicadas", ext="xlsx"),
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    key=f"{prefix}btn_duplicadas_xlsx",
                                )

                        extrair_sem_cpf = config.get("extrair_sem_cpf")
                        if extrair_sem_cpf is not None and df_aluno_cur is not None:
                            sem_cpf = extrair_sem_cpf(df_aluno_cur)
                            if sem_cpf:
                                st.download_button(
                                    "Matrículas sem CPF informado (XLSX)",
                                    data=_sem_cpf_xlsx(sem_cpf),
                                    file_name=_output_filename(config, suffix="alunos_sem_cpf", ext="xlsx"),
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    key=f"{prefix}btn_sem_cpf_xlsx",
                                )

                        levantamento_cor_raca = config.get("levantamento_cor_raca")
                        if levantamento_cor_raca is not None and df_aluno_cur is not None:
                            levantamento = levantamento_cor_raca(df_aluno_cur)
                            st.download_button(
                                "Quantitativo de Cor/Raça por Unidade Escolar",
                                data=_levantamento_cor_raca_xlsx(levantamento),
                                file_name=_output_filename(config, suffix="levantamento_cor_raca", ext="xlsx"),
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"{prefix}btn_cor_raca_xlsx",
                            )

                    if resultados is None:
                        df_cur = st.session_state.get(f"{prefix}df_aluno_curricular")
                        ano_letivo = st.session_state.get(f"{prefix}ano_letivo")
                        with st.spinner("Carregando Relatório de Inconsistências..."):
                            resultados = config["validate"](df_cur, escolas, ano_letivo)
                        st.session_state[f"{prefix}resultados"] = resultados
                        st.session_state[f"{prefix}finalizado"] = True
                        finalizado = True

                if not is_aluno:
                    st.divider()

                    if st.button("Iniciar Verificação de Inconsistências", key=f"{prefix}btn_validar", type="primary"):
                        with st.spinner("Carregando Relatório de Inconsistências..."):
                            if is_turmas:
                                df_aee = st.session_state.get(f"{prefix}df_aee")
                                df_atv = st.session_state.get(f"{prefix}df_atv")
                                df_cur = st.session_state.get(f"{prefix}df_curricular")
                                df_escol = st.session_state.get(f"{prefix}df_escolarizacao")
                                resultados = config["validate"](df_aee, df_atv, df_cur, df_escol)
                            elif is_prof:
                                df_prof_cur = st.session_state.get(f"{prefix}df_prof_curricular")
                                resultados = config["validate"](df_prof_cur, escolas)
                            elif is_zip:
                                df_unico = st.session_state.get(f"{prefix}df_unificado")
                                resultados = config["validate"](df_unico, escolas)
                            else:
                                raw = st.session_state.get(f"{prefix}arquivo_bytes")
                                inicio = st.session_state.get(f"{prefix}inicio_ano_letivo")
                                termino = st.session_state.get(f"{prefix}termino_ano_letivo")
                                resultados = config["validate"](raw, escolas, inicio, termino)
                        st.session_state[f"{prefix}resultados"] = resultados
                        st.session_state[f"{prefix}finalizado"] = True
                        st.rerun()
        else:
            st.success("Estrutura da planilha verificada.")
            if st.button("Iniciar validacoes", key=f"{prefix}btn_validar", type="primary"):
                raw = st.session_state.get(f"{prefix}arquivo_bytes")
                inicio = st.session_state.get(f"{prefix}inicio_ano_letivo")
                termino = st.session_state.get(f"{prefix}termino_ano_letivo")
                with st.spinner("Carregando Relatório de Inconsistências..."):
                    resultados = config["validate"](raw, escolas, inicio, termino)
                st.session_state[f"{prefix}resultados"] = resultados
                st.session_state[f"{prefix}finalizado"] = True
                st.rerun()

    # Etapa 3: resultados prontos -> exibir resumo + download
    if resultados is not None and finalizado:
        incon = [r for r in resultados if r.get("status") == "inconsistente"]

        if incon:
            if not is_aluno:
                unids = len({r.get('codigo_inep', '') for r in incon})
                st.markdown(f"**Total de Inconsistencias:** {len(incon)}")
                st.markdown(f"**Total de Unidades:** {unids}")

            with st.spinner("Carregando Relatório de Inconsistências..."):
                xlsx_bytes = _build_xlsx(resultados, modo=_modo_relatorio(config))
            st.download_button(
                label="Relatório de Inconsistências (XLSX)",
                data=xlsx_bytes,
                file_name=_output_filename(config),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    # Etapa 1: formulario submetido -> processar
    if not submitted:
        return

    if not arquivo:
        st.error("Envie o arquivo de dados para verificacao.")
        return

    raw_bytes = arquivo.read()

    if is_aluno:
        st.session_state[f"{prefix}ano_letivo"] = ano_letivo

    if config.get("requires_dates", True):
        st.session_state[f"{prefix}inicio_ano_letivo"] = inicio
        st.session_state[f"{prefix}termino_ano_letivo"] = termino

    if is_zip:
        try:
            if is_turmas:
                dados = read_turmas_zip(raw_bytes)
            elif is_prof:
                dados = read_profissionais_zip(raw_bytes)
            elif is_aluno:
                dados = read_aluno_zip(raw_bytes)
            else:
                dados = read_zip_merged(raw_bytes)
        except Exception as exc:
            st.error(f"Falha ao processar o ZIP: {exc}")
            return

        escolas = dados["escolas"]

        if not escolas:
            st.error("Nenhuma escola encontrada nos CSVs do ZIP.")
            return

        st.session_state[f"{prefix}df_unificado"] = dados["unificado"]
        if is_turmas:
            st.session_state[f"{prefix}df_aee"] = dados["turmas_aee"]
            st.session_state[f"{prefix}df_atv"] = dados["turmas_atvcomplementar"]
            st.session_state[f"{prefix}df_curricular"] = dados["turmas_curricular"]
            st.session_state[f"{prefix}df_escolarizacao"] = dados["turmas_escolarizacao"]
        if is_prof:
            st.session_state[f"{prefix}df_prof_aee_atv"] = dados["prof_aee_atv"]
            st.session_state[f"{prefix}df_prof_curricular"] = dados["prof_curricular"]
        if is_aluno:
            st.session_state[f"{prefix}df_aluno_curricular"] = dados["aluno_curricular"]
            st.session_state[f"{prefix}df_aluno_aee"] = dados["aluno_aee"]
            st.session_state[f"{prefix}df_aluno_atv"] = dados["aluno_atv"]
            st.session_state[f"{prefix}df_aluno_transporte"] = dados["aluno_transporte"]
            st.session_state[f"{prefix}df_aluno_deficiencia"] = dados["aluno_deficiencia"]
        st.session_state[f"{prefix}escolas"] = escolas
        st.session_state[f"{prefix}estrutura_ok"] = True
        st.session_state[f"{prefix}arquivo_bytes"] = raw_bytes
        st.session_state[f"{prefix}resultados"] = None
        st.session_state[f"{prefix}finalizado"] = False
        esperadas = config.get("expected_columns") or []
        ok_col, faltando, extras = checar_colunas(dados["unificado"], esperadas)
        st.session_state[f"{prefix}colunas_validadas"] = ok_col
        st.session_state[f"{prefix}colunas_erros"] = (faltando, extras) if not ok_col else None

        st.success(
            f"{dados['csv_count']} arquivo(s) CSV descompactado(s) do ZIP. "
            f"{len(dados['unificado'])} registro(s) totais."
        )
        st.rerun()

    ref_bytes = _REFERENCIA_BYTES.get(config["key"])
    if ref_bytes is not None:
        valido, erros = checar_estrutura(raw_bytes, ref_bytes)
        if not valido:
            st.error("A estrutura do arquivo nao corresponde ao modelo esperado:")
            for e in erros:
                st.markdown(f"- {e}")
            return

    try:
        escolas = read_school_list(io.BytesIO(raw_bytes))
    except Exception as exc:
        st.error(f"Falha ao ler o arquivo: {exc}")
        return

    if not escolas:
        st.error("Nenhum registro valido encontrado no arquivo.")
        return

    st.session_state[f"{prefix}escolas"] = escolas
    st.session_state[f"{prefix}estrutura_ok"] = True
    st.session_state[f"{prefix}arquivo_bytes"] = raw_bytes
    st.session_state[f"{prefix}resultados"] = None
    st.session_state[f"{prefix}finalizado"] = False

    st.rerun()