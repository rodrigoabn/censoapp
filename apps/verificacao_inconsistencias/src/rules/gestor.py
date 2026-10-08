from __future__ import annotations

import io
from typing import Any

import openpyxl

SHEET_VINCULO = "Vínculo do gestor"
SHEET_IDENT = "Identificação"
SHEET_DADOS = "Dados pessoais"


def _val(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def validate_gestor(
    raw_bytes: bytes,
    escolas: list[dict],
    inicio_ano_letivo=None,
    termino_ano_letivo=None,
) -> list[dict]:
    resultados: list[dict] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception as exc:
        return [{
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "gestor",
            "status": "erro",
            "detalhes": f"Não foi possível ler o arquivo: {exc}",
        }]

    # ═════════════════════════════════════════════════════════════════
    # Validação da aba "Vínculo do gestor"
    # ═════════════════════════════════════════════════════════════════

    if SHEET_VINCULO not in wb.sheetnames:
        return [{
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "gestor",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_VINCULO}' não encontrada no arquivo.",
        }]

    ws = wb[SHEET_VINCULO]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    if rows:
        headers = [str(c).strip() if c else "" for c in next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True)
        )]

        def col_idx(name: str) -> int:
            for i, h in enumerate(headers):
                if h == name:
                    return i
            return -1

        col_cargo = col_idx("1 – Cargo")
        col_acesso = col_idx("2 - Critério de acesso ao cargo/função")
        col_vinculo = col_idx("3 - Situação Funcional/Regime de contratação/Tipo de vínculo")
        col_email = col_idx("Email principal")

        _VALORES_PROIBIDOS_ACESSO = {
            "Exclusivamente por processo eleitoral com a participação da comunidade escolar",
            "Concurso público específico para o cargo de gestor escolar",
        }

        for row in rows:
            vals = [_val(c) for c in row]
            codigo = vals[0] if len(vals) > 0 else ""
            nome = vals[1] if len(vals) > 1 else ""

            # Regra: "1 – Cargo" = "Diretor(a)"
            if col_cargo >= 0 and len(vals) > col_cargo:
                v = vals[col_cargo]
                if v.lower() != "diretor(a)":
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_VINCULO,
                        "status": "inconsistente",
                        "detalhes": f"Coluna '1 – Cargo': esperado 'Diretor(a)', encontrado '{v}'",
                    })

            # Regra: "2 - Critério de acesso..." não pode ser proibido nem vazio
            if col_acesso >= 0 and len(vals) > col_acesso:
                v = vals[col_acesso]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_VINCULO,
                        "status": "inconsistente",
                        "detalhes": "Coluna '2 - Critério de acesso ao cargo/função': não pode estar vazia",
                    })
                elif v in _VALORES_PROIBIDOS_ACESSO:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_VINCULO,
                        "status": "inconsistente",
                        "detalhes": f"Coluna '2 - Critério de acesso ao cargo/função': valor não permitido '{v}'",
                    })

            # Regra: "3 - Situação Funcional..." ∈ {"Concursado/efetivo/estável", "Contrato temporário"}
            if col_vinculo >= 0 and len(vals) > col_vinculo:
                v = vals[col_vinculo]
                if v.lower() not in ("concursado/efetivo/estável", "contrato temporário"):
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_VINCULO,
                        "status": "inconsistente",
                        "detalhes": (
                            f"Coluna '3 - Situação Funcional/Regime de contratação/"
                            f"Tipo de vínculo': esperado 'Concursado/efetivo/estável' "
                            f"ou 'Contrato temporário', encontrado '{v}'"
                        ),
                    })

            # Regra: "Email principal" não vazio
            if col_email >= 0 and len(vals) > col_email:
                v = vals[col_email]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_VINCULO,
                        "status": "inconsistente",
                        "detalhes": "Coluna 'Email principal': não pode estar vazia",
                    })

    # ═════════════════════════════════════════════════════════════════
    # Validação da aba "Identificação"
    # ═════════════════════════════════════════════════════════════════

    if SHEET_IDENT not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": SHEET_IDENT,
            "status": "erro",
            "detalhes": f"Aba '{SHEET_IDENT}' não encontrada no arquivo.",
        })
        return resultados

    ws2 = wb[SHEET_IDENT]
    rows2 = list(ws2.iter_rows(min_row=2, values_only=True))

    if rows2:
        headers2 = [str(c).strip() if c else "" for c in next(
            ws2.iter_rows(min_row=1, max_row=1, values_only=True)
        )]

        def col_idx2(name: str) -> int:
            for i, h in enumerate(headers2):
                if h == name:
                    return i
            return -1

        col_filiacao1 = col_idx2("5a - Nome completo da filiação 1")
        col_filiacao2 = col_idx2("5b - Nome completo da filiação 2")
        col_sexo = col_idx2("6 – Sexo")
        col_cor = col_idx2("7 – Cor/Raça")
        col_nacionalidade = col_idx2("8 – Nacionalidade")
        col_uf_nasc = col_idx2("10 - UF de nascimento")
        col_municipio_nasc = col_idx2("11 - Município de nascimento")

        for row in rows2:
            vals = [_val(c) for c in row]
            codigo = vals[0] if len(vals) > 0 else ""
            nome = vals[1] if len(vals) > 1 else ""

            if col_filiacao1 >= 0 and len(vals) > col_filiacao1:
                v = vals[col_filiacao1]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '5a - Nome completo da filiação 1': não pode estar vazia",
                    })

            if col_filiacao2 >= 0 and len(vals) > col_filiacao2:
                v = vals[col_filiacao2]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '5b - Nome completo da filiação 2': não pode estar vazia",
                    })

            if col_sexo >= 0 and len(vals) > col_sexo:
                v = vals[col_sexo]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '6 – Sexo': não pode estar vazia",
                    })

            if col_cor >= 0 and len(vals) > col_cor:
                v = vals[col_cor]
                if v.lower() in ("não declarada", "nao declarada"):
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": f"Coluna '7 – Cor/Raça': não pode ser 'Não declarada', encontrado '{v}'",
                    })

            if col_nacionalidade >= 0 and len(vals) > col_nacionalidade:
                v = vals[col_nacionalidade]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '8 – Nacionalidade': não pode estar vazia",
                    })

            if col_uf_nasc >= 0 and len(vals) > col_uf_nasc:
                v = vals[col_uf_nasc]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '10 - UF de nascimento': não pode estar vazia",
                    })

            if col_municipio_nasc >= 0 and len(vals) > col_municipio_nasc:
                v = vals[col_municipio_nasc]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_IDENT,
                        "status": "inconsistente",
                        "detalhes": "Coluna '11 - Município de nascimento': não pode estar vazia",
                    })

    # ═════════════════════════════════════════════════════════════════
    # Validação da aba "Dados pessoais"
    # ═════════════════════════════════════════════════════════════════

    if SHEET_DADOS not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": SHEET_DADOS,
            "status": "erro",
            "detalhes": f"Aba '{SHEET_DADOS}' não encontrada no arquivo.",
        })
        return resultados

    ws3 = wb[SHEET_DADOS]
    rows3 = list(ws3.iter_rows(min_row=2, values_only=True))

    if rows3:
        headers3 = [str(c).strip() if c else "" for c in next(
            ws3.iter_rows(min_row=1, max_row=1, values_only=True)
        )]

        def col_idx3(name: str) -> int:
            for i, h in enumerate(headers3):
                if h == name:
                    return i
            return -1

        col_pais = col_idx3("13 - País de residência")
        col_cep = col_idx3("14 – CEP")
        col_uf = col_idx3("15 – UF")
        col_municipio = col_idx3("16 – Município")
        col_zona = col_idx3("17 - Localização/zona de residência")
        col_localizacao_dif = col_idx3("18 - Localização diferenciada da residência")
        col_escolaridade = col_idx3("19 - Maior nível de escolaridade concluído")
        col_ensino_medio = col_idx3("19a - Tipo de ensino médio cursado")
        col_pos = col_idx3("20. Pós-graduações concluídas")

        for row in rows3:
            vals = [_val(c) for c in row]
            codigo = vals[0] if len(vals) > 0 else ""
            nome = vals[1] if len(vals) > 1 else ""

            if col_pais >= 0 and len(vals) > col_pais:
                v = vals[col_pais]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '13 - País de residência': não pode estar vazia",
                    })

            if col_cep >= 0 and len(vals) > col_cep:
                v = vals[col_cep]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '14 – CEP': não pode estar vazia",
                    })

            if col_uf >= 0 and len(vals) > col_uf:
                v = vals[col_uf]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '15 – UF': não pode estar vazia",
                    })

            if col_municipio >= 0 and len(vals) > col_municipio:
                v = vals[col_municipio]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '16 – Município': não pode estar vazia",
                    })

            if col_zona >= 0 and len(vals) > col_zona:
                v = vals[col_zona]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '17 - Localização/zona de residência': não pode estar vazia",
                    })

            if col_localizacao_dif >= 0 and len(vals) > col_localizacao_dif:
                v = vals[col_localizacao_dif]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '18 - Localização diferenciada da residência': não pode estar vazia",
                    })

            if col_escolaridade >= 0 and len(vals) > col_escolaridade:
                v = vals[col_escolaridade]
                if v.lower() != "educação superior":
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": (
                            f"Coluna '19 - Maior nível de escolaridade concluído': "
                            f"esperado 'Educação superior', encontrado '{v}'"
                        ),
                    })

            if col_ensino_medio >= 0 and len(vals) > col_ensino_medio:
                v = vals[col_ensino_medio]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '19a - Tipo de ensino médio cursado': não pode estar vazia",
                    })

            if col_pos >= 0 and len(vals) > col_pos:
                v = vals[col_pos]
                if not v:
                    resultados.append({
                        "codigo_inep": codigo,
                        "nome_unidade": nome,
                        "tipo_validacao": SHEET_DADOS,
                        "status": "inconsistente",
                        "detalhes": "Coluna '20. Pós-graduações concluídas': não pode estar vazia",
                    })

    return resultados
