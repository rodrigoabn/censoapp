from __future__ import annotations

import io
from typing import Any

import openpyxl

SHEET_NAME = "Vinculação institucional e Conv"
SHEET_NAME_FUNC = "Funcionamento e Identificação"
SHEET_NAME_ESTRUTURA = "Estrutura física"
SHEET_NAME_EQUIP = "Equipamentos e recursos tecnoló"
SHEET_NAME_ORG = "Organização escolar"


def _val(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def validate_unidades(
    raw_bytes: bytes,
    escolas: list[dict],
    inicio_ano_letivo=None,
    termino_ano_letivo=None,
) -> list[dict]:
    resultados: list[dict] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
    except Exception as exc:
        return [{
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "unidades",
            "status": "erro",
            "detalhes": f"Não foi possível ler o arquivo: {exc}",
        }]

    if SHEET_NAME not in wb.sheetnames:
        return [{
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "unidades",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_NAME}' não encontrada no arquivo.",
        }]

    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    if not rows:
        return resultados

    headers = [str(c).strip() if c else "" for c in next(
        ws.iter_rows(min_row=1, max_row=1, values_only=True)
    )]

    def col_idx(name: str) -> int:
        for i, h in enumerate(headers):
            if h == name:
                return i
        return -1

    col_regulamentacao = col_idx(
        "2 - Regulamentação/autorização no conselho ou órgão municipal, estadual ou federal de educação"
    )
    col_estadual = col_idx("Estadual")
    col_municipal = col_idx("Municipal")
    col_localizacao = col_idx("3 – localizacaoZona")
    col_parceria = col_idx(
        "5 - A escola possui parceria ou convênio com a Administração Pública e/ou outras instituições"
    )

    for row_idx, row in enumerate(rows, start=2):
        vals = [_val(c) for c in row]
        codigo = vals[0] if len(vals) > 0 else ""
        nome = vals[1] if len(vals) > 1 else ""

        # Regra 1: "2 - Regulamentação..." deve ser "Sim"
        if col_regulamentacao >= 0 and len(vals) > col_regulamentacao:
            v = vals[col_regulamentacao]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Vinculação institucional e Conv",
                    "status": "inconsistente",
                    "detalhes": (
                        f"Coluna '2 - Regulamentação/autorização no conselho ou órgão "
                        f"municipal, estadual ou federal de educação': esperado 'Sim', "
                        f"encontrado '{v}'"
                    ),
                })

        # Regra 2: "Estadual" deve ser vazia
        if col_estadual >= 0 and len(vals) > col_estadual:
            v = vals[col_estadual]
            if v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Vinculação institucional e Conv",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Estadual': esperado vazio, encontrado '{v}'",
                })

        # Regra 3: "Municipal" deve ser "Sim"
        if col_municipal >= 0 and len(vals) > col_municipal:
            v = vals[col_municipal]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Vinculação institucional e Conv",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Municipal': esperado 'Sim', encontrado '{v}'",
                })

        # Regra 4: "3 – localizacaoZona" nunca vazia
        if col_localizacao >= 0 and len(vals) > col_localizacao:
            v = vals[col_localizacao]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Vinculação institucional e Conv",
                    "status": "inconsistente",
                    "detalhes": "Coluna '3 – localizacaoZona': não pode estar vazia",
                })

        # Regra 5: "5 - A escola possui parceria..." deve ser "Não"
        if col_parceria >= 0 and len(vals) > col_parceria:
            v = vals[col_parceria]
            if v.lower() not in ("não", "nao"):
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Vinculação institucional e Conv",
                    "status": "inconsistente",
                    "detalhes": (
                        f"Coluna '5 - A escola possui parceria ou convênio com a "
                        f"Administração Pública e/ou outras instituições': esperado "
                        f"'Não', encontrado '{v}'"
                    ),
                })

    # ═════════════════════════════════════════════════════════════════
    # Validação da aba "Funcionamento e Identificação"
    # ═════════════════════════════════════════════════════════════════

    _date_fmt = "%d/%m/%Y"
    _inicio_str = None
    _termino_str = None
    if inicio_ano_letivo:
        _inicio_str = inicio_ano_letivo.strftime(_date_fmt)
    if termino_ano_letivo:
        _termino_str = termino_ano_letivo.strftime(_date_fmt)

    if SHEET_NAME_FUNC not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "Funcionamento e Identificação",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_NAME_FUNC}' não encontrada no arquivo.",
        })
        return resultados

    ws2 = wb[SHEET_NAME_FUNC]
    rows2 = list(ws2.iter_rows(min_row=2, values_only=True))

    if not rows2:
        return resultados

    headers2 = [str(c).strip() if c else "" for c in next(
        ws2.iter_rows(min_row=1, max_row=1, values_only=True)
    )]

    def col_idx2(name: str) -> int:
        for i, h in enumerate(headers2):
            if h == name:
                return i
        return -1

    col_situacao = col_idx2("6 - Situação de funcionamento")
    col_7a = col_idx2("7a – Início")
    col_7b = col_idx2("7b -Término (previsão)")
    col_cep = col_idx2("9 - CEP")
    col_distrito = col_idx2("11b – Distrito")
    col_email = col_idx2("19 - Endereço eletrônico (e-mail) da escola")
    col_localizacao_dif = col_idx2("20 - Localização diferenciada da escola")

    for row in rows2:
        vals = [_val(c) for c in row]
        codigo = vals[0] if len(vals) > 0 else ""
        nome = vals[1] if len(vals) > 1 else ""

        if col_situacao >= 0 and len(vals) > col_situacao:
            v = vals[col_situacao]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": "Coluna '6 - Situação de funcionamento': não pode estar vazia",
                })

        if col_7a >= 0 and len(vals) > col_7a and _inicio_str:
            v = vals[col_7a]
            if v != _inicio_str:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": f"Coluna '7a – Início': esperado '{_inicio_str}', encontrado '{v}'",
                })

        if col_7b >= 0 and len(vals) > col_7b and _termino_str:
            v = vals[col_7b]
            if v != _termino_str:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": f"Coluna '7b -Término (previsão)': esperado '{_termino_str}', encontrado '{v}'",
                })

        if col_cep >= 0 and len(vals) > col_cep:
            v = vals[col_cep]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": "Coluna '9 - CEP': não pode estar vazia",
                })

        if col_distrito >= 0 and len(vals) > col_distrito:
            v = vals[col_distrito]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": "Coluna '11b – Distrito': não pode estar vazia",
                })

        if col_email >= 0 and len(vals) > col_email:
            v = vals[col_email]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": "Coluna '19 - Endereço eletrônico (e-mail) da escola': não pode estar vazia",
                })

        if col_localizacao_dif >= 0 and len(vals) > col_localizacao_dif:
            v = vals[col_localizacao_dif]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Funcionamento e Identificação",
                    "status": "inconsistente",
                    "detalhes": "Coluna '20 - Localização diferenciada da escola': não pode estar vazia",
                })


    # ================================================================
    # Validação da aba "Estrutura física"
    # ================================================================

    if SHEET_NAME_ESTRUTURA not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "Estrutura física",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_NAME_ESTRUTURA}' não encontrada no arquivo.",
        })
        return resultados

    ws3 = wb[SHEET_NAME_ESTRUTURA]
    rows3 = list(ws3.iter_rows(min_row=2, values_only=True))

    if not rows3:
        return resultados

    headers3 = [str(c).strip() if c else "" for c in next(
        ws3.iter_rows(min_row=1, max_row=1, values_only=True)
    )]

    def col_idx3(name: str) -> int:
        for i, h in enumerate(headers3):
            if h == name:
                return i
        return -1

    col_predio = col_idx3("Prédio escolar")
    col_ocupacao = col_idx3("28 - Forma de ocupação do prédio escolar")
    col_compartilha = col_idx3("29 - A escola compartilha o seu prédio com outra instituição de ensino")
    col_agua = col_idx3("30 - Fornece água potável para o consumo humano")
    col_dormitorio = col_idx3("Dormitório de professor(a)")
    col_lab_prof = col_idx3("Laboratório específico para a educação profissional")
    col_oficina = col_idx3("Sala de oficinas da educação profissional")

    for row in rows3:
        vals = [_val(c) for c in row]
        codigo = vals[0] if len(vals) > 0 else ""
        nome = vals[1] if len(vals) > 1 else ""

        if col_predio >= 0 and len(vals) > col_predio:
            v = vals[col_predio]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Prédio escolar': esperado 'Sim', encontrado '{v}'",
                })

        if col_ocupacao >= 0 and len(vals) > col_ocupacao:
            v = vals[col_ocupacao]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": "Coluna '28 - Forma de ocupação do prédio escolar': não pode estar vazia",
                })

        if col_compartilha >= 0 and len(vals) > col_compartilha:
            v = vals[col_compartilha]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": "Coluna '29 - A escola compartilha o seu prédio com outra instituição de ensino': não pode estar vazia",
                })

        if col_agua >= 0 and len(vals) > col_agua:
            v = vals[col_agua]
            if not v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": "Coluna '30 - Fornece água potável para o consumo humano': não pode estar vazia",
                })

        if col_dormitorio >= 0 and len(vals) > col_dormitorio:
            v = vals[col_dormitorio]
            if v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Dormitório de professor(a)': esperado vazio, encontrado '{v}'",
                })

        if col_lab_prof >= 0 and len(vals) > col_lab_prof:
            v = vals[col_lab_prof]
            if v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Laboratório específico para a educação profissional': esperado vazio, encontrado '{v}'",
                })

        if col_oficina >= 0 and len(vals) > col_oficina:
            v = vals[col_oficina]
            if v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Estrutura física",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Sala de oficinas da educação profissional': esperado vazio, encontrado '{v}'",
                })


    # ================================================================
    # Validação da aba "Equipamentos e recursos tecnoló"
    # ================================================================

    if SHEET_NAME_EQUIP not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "Equipamentos e recursos tecnoló",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_NAME_EQUIP}' nao encontrada no arquivo.",
        })
        return resultados

    ws4 = wb[SHEET_NAME_EQUIP]
    rows4 = list(ws4.iter_rows(min_row=2, values_only=True))

    if not rows4:
        return resultados

    headers4 = [str(c).strip() if c else "" for c in next(
        ws4.iter_rows(min_row=1, max_row=1, values_only=True)
    )]

    def col_idx4(name: str) -> int:
        for i, h in enumerate(headers4):
            if h == name:
                return i
        return -1

    col_computadores = col_idx4("Computadores")
    col_sem_internet = col_idx4("Não possui acesso à internet")
    col_uso_admin = col_idx4("Para uso administrativo")
    col_banda_larga = col_idx4("47 - Internet banda larga")

    for row in rows4:
        vals = [_val(c) for c in row]
        codigo = vals[0] if len(vals) > 0 else ""
        nome = vals[1] if len(vals) > 1 else ""

        if col_computadores >= 0 and len(vals) > col_computadores:
            v = vals[col_computadores]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Equipamentos e recursos tecnoló",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Computadores': esperado 'Sim', encontrado '{v}'",
                })

        if col_sem_internet >= 0 and len(vals) > col_sem_internet:
            v = vals[col_sem_internet]
            if v:
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Equipamentos e recursos tecnoló",
                    "status": "inconsistente",
                    "detalhes": "Coluna 'Não possui acesso à internet': deve estar vazia",
                })

        if col_uso_admin >= 0 and len(vals) > col_uso_admin:
            v = vals[col_uso_admin]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Equipamentos e recursos tecnoló",
                    "status": "inconsistente",
                    "detalhes": f"Coluna 'Para uso administrativo': esperado 'Sim', encontrado '{v}'",
                })

        if col_banda_larga >= 0 and len(vals) > col_banda_larga:
            v = vals[col_banda_larga]
            if v.lower() != "sim":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Equipamentos e recursos tecnoló",
                    "status": "inconsistente",
                    "detalhes": f"Coluna '47 - Internet banda larga': esperado 'Sim', encontrado '{v}'",
                })

    # ================================================================
    # Validação da aba "Organização escolar"
    # ================================================================

    if SHEET_NAME_ORG not in wb.sheetnames:
        resultados.append({
            "codigo_inep": "",
            "nome_unidade": "",
            "tipo_validacao": "Organização escolar",
            "status": "erro",
            "detalhes": f"Aba '{SHEET_NAME_ORG}' não encontrada no arquivo.",
        })
        return resultados

    ws5 = wb[SHEET_NAME_ORG]
    rows5 = list(ws5.iter_rows(min_row=2, values_only=True))

    if not rows5:
        return resultados

    headers5 = [str(c).strip() if c else "" for c in next(
        ws5.iter_rows(min_row=1, max_row=1, values_only=True)
    )]

    def col_idx5(name: str) -> int:
        for i, h in enumerate(headers5):
            if h == name:
                return i
        return -1

    col_alimentacao = col_idx5(
        "49 - A escola fornece alimentação escolar para os alunos"
    )
    col_indigena = col_idx5("50 - Escola indígena *")
    col_exame = col_idx5(
        "58 - A escola faz exame de seleção para ingresso de seus alunos "
        "(avaliação por prova e/ou análise curricular)"
    )

    for row in rows5:
        vals = [_val(c) for c in row]
        codigo = vals[0] if len(vals) > 0 else ""
        nome = vals[1] if len(vals) > 1 else ""

        if col_alimentacao >= 0 and len(vals) > col_alimentacao:
            v = vals[col_alimentacao]
            if v.lower() != "oferece":
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Organização escolar",
                    "status": "inconsistente",
                    "detalhes": (
                        f"Coluna '49 - A escola fornece alimentação escolar "
                        f"para os alunos': diferentes de 'Oferece', "
                        f"encontrado '{v}'"
                    ),
                })

        if col_indigena >= 0 and len(vals) > col_indigena:
            v = vals[col_indigena]
            if v.lower() not in ("não", "nao"):
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Organização escolar",
                    "status": "inconsistente",
                    "detalhes": (
                        f"Coluna '50 - Escola indígena *': esperado 'Não', "
                        f"encontrado '{v}'"
                    ),
                })

        if col_exame >= 0 and len(vals) > col_exame:
            v = vals[col_exame]
            if v.lower() not in ("não", "nao"):
                resultados.append({
                    "codigo_inep": codigo,
                    "nome_unidade": nome,
                    "tipo_validacao": "Organização escolar",
                    "status": "inconsistente",
                    "detalhes": (
                        f"Coluna '58 - A escola faz exame de seleção para "
                        f"ingresso de seus alunos (avaliação por prova e/ou "
                        f"análise curricular)': esperado 'Não', encontrado '{v}'"
                    ),
                })

    return resultados
